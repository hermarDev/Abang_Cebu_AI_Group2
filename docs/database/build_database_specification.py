#!/usr/bin/env python3
"""Build SCRUM-146 workbook from catalog metadata in the local applied Supabase DB."""
from __future__ import annotations
import json, subprocess, hashlib, os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parent
DB=os.environ.get('SUPABASE_DB_CONTAINER','supabase_db_scrum-143-search-db')  # Local inventory only; never use a production target.
TEMPLATE=ROOT/'docs/templates/Database_Specification_Template.xlsx'
OUTPUT=ROOT/'docs/database/Database_Specification.xlsx'
TABLES=['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']


def q(sql):
    r=subprocess.run(['docker','exec',DB,'psql','-U','postgres','-d','postgres','-X','-qAt','-c',sql],capture_output=True,text=True)
    if r.returncode: raise RuntimeError(r.stderr)
    return json.loads(r.stdout.strip() or 'null')

def rows(sql): return q(f"select coalesce(json_agg(row_to_json(q)), '[]'::json) from ({sql}) q")
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()

cols=rows("""select table_name,column_name,ordinal_position,data_type,udt_name,is_nullable,column_default,
character_maximum_length,numeric_precision,numeric_scale,datetime_precision,
col_description((quote_ident(table_schema)||'.'||quote_ident(table_name))::regclass::oid,ordinal_position) as description
from information_schema.columns where table_schema='public' and table_name=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by table_name,ordinal_position""")
constraints=rows("""select c.relname as table_name, con.conname as name, con.contype::text as kind, pg_get_constraintdef(con.oid,true) as definition
from pg_constraint con join pg_class c on c.oid=con.conrelid join pg_namespace n on n.oid=c.relnamespace
where n.nspname='public' and c.relname=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by c.relname,con.conname""")
indexes=rows("""select tablename as table_name,indexname as name,indexdef as definition from pg_indexes
where schemaname='public' and tablename=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by tablename,indexname""")
policies=rows("""select tablename as table_name,policyname as name,cmd as command,roles,coalesce(qual,'') as using_expression,coalesce(with_check,'') as check_expression
from pg_policies where schemaname='public' and tablename=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by tablename,policyname""")
triggers=rows("""select event_object_table as table_name,trigger_name as name,event_manipulation as event,action_timing as timing,action_orientation as orientation,action_statement as definition
from information_schema.triggers where trigger_schema='public' and event_object_table=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by event_object_table,trigger_name,event_manipulation""")
rls=rows("""select c.relname as table_name,c.relrowsecurity as enabled,c.relforcerowsecurity as forced from pg_class c join pg_namespace n on n.oid=c.relnamespace
where n.nspname='public' and c.relname=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by c.relname""")
enums=rows("""select t.typname as name,'public.'||t.typname as pg_type,array_agg(e.enumlabel order by e.enumsortorder) as values
from pg_type t join pg_enum e on e.enumtypid=t.oid join pg_namespace n on n.oid=t.typnamespace where n.nspname='public' group by t.typname order by t.typname""")
migrations=rows("""select version,name,cardinality(statements) as statement_count from supabase_migrations.schema_migrations order by version""")
functions=rows("""select p.proname as name,pg_get_function_identity_arguments(p.oid) as arguments,pg_get_function_result(p.oid) as returns,
case when p.prosecdef then 'DEFINER' else 'INVOKER' end as security,
coalesce(array_to_string(p.proconfig,', '),'') as configuration,
coalesce(obj_description(p.oid,'pg_proc'),'') as description
from pg_proc p join pg_namespace n on n.oid=p.pronamespace where n.nspname='public' and p.prokind in ('f','p') and p.proname in ('handle_updated_at','get_current_user_role','is_admin','handle_new_user','prevent_audit_log_mutation','search_listings_in_viewport','search_listings_near_landmark') order by p.proname""")
tables=rows("""select c.relname as name,coalesce(obj_description(c.oid,'pg_class'),'') as description from pg_class c join pg_namespace n on n.oid=c.relnamespace
where n.nspname='public' and c.relkind='r' and c.relname=any(ARRAY['profiles','kyc_verifications','properties','rental_units','property_photos','amenities','property_amenities','inquiries','saved_listings','reviews','audit_logs']) order by c.relname""")
assert len(cols)==100 and len(tables)==11 and len(indexes)==39 and len(constraints)==67 and len(migrations)==5
assert {x['name'] for x in tables}==set(TABLES)

migration_sources=[
 ('20260929000000_init_extensions','origin/main @ be77fa4aedd8f352614c731de48f72288a603a75','supabase/migrations/20260929000000_init_extensions.sql','b75d7a00016c3b69985e608992ab8e369ccad7742cc9874a08391dfbd4d474ed'),
 ('20260929000001_users_and_profiles','origin/main @ be77fa4aedd8f352614c731de48f72288a603a75','supabase/migrations/20260929000001_users_and_profiles.sql','248a9ad9cd0c748ca0c2d4c57628b0ea86036aa8a1bd30e68d755acb0e01b88b'),
 ('20261019000140_listings','PR #7 head @ 25fe87492b164438b29bb9f3b82f42b9770a43af','supabase/migrations/20261019000140_listings.sql','14840b35f23c5f185e8056c83595369e9e325b8f3839df5d6d42884ecab5b514'),
 ('20261019000141_engagement','PR #7 head @ 25fe87492b164438b29bb9f3b82f42b9770a43af','supabase/migrations/20261019000141_engagement.sql','35ee513041bb1cd70e6558225422f76eee86e1b84a7d1befe81aefd249b9a884'),
 ('20261019000143_search_functions','PR #6 head @ c64e8a08a247799131597d5f39f21a1aee79fa50','supabase/migrations/20261019000143_search_functions.sql','fdcb8602bf705a10df34421c1006505a103ebd01c6642f2cc9715d0f32753788')]

wb=load_workbook(TEMPLATE)
wb.remove(wb['<TABLE_NAME>'])
rev=wb['Revision_History']; rev.delete_rows(3,rev.max_row)
rev['B3']='2.0'; rev['C3']='2026-10-21'; rev['D3']='Hermar Centillas'; rev['E3']='Completed applied-schema inventory: 11 public tables, 100 columns, constraints, indexes, relationships and migration provenance.'
listing=wb['List of Tables']; listing.delete_rows(3,listing.max_row); listing['B2']='No.';listing['C2']='Table Name';listing['D2']='Description'
for i,t in enumerate(tables,1):
 r=i+2; listing.cell(r,2,i); listing.cell(r,3,'public.'+t['name']); listing.cell(r,4,t['description'] or 'No PostgreSQL table comment recorded.')
 for c in range(2,5): listing.cell(r,c).alignment=Alignment(vertical='top',wrap_text=True)
listing.freeze_panes='B3'; listing.auto_filter.ref='B2:D13'
for col,width in {'A':6,'B':9,'C':31,'D':76}.items(): listing.column_dimensions[col].width=width

for name in TABLES:
 ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
 ws['B2']='Table Name';ws['C2']='public.'+name
 ws['B3']='Updated By';ws['C3']='Hermar Centillas'
 ws['B4']='Date of Update';ws['C4']='2026-10-21'
 ws['A6']='NO.';ws['B6']='Column Name';ws['C6']='Data Type';ws['D6']='Size/Length';ws['E6']='Nullable (Yes/No)';ws['F6']='Primary Key (Yes/No)';ws['G6']='Default Value';ws['H6']='NOTES/CONSTRAINTS'
 for cell in ws[6]:
  cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78');cell.alignment=Alignment(vertical='center',horizontal='center',wrap_text=True)
 table_cols=[x for x in cols if x['table_name']==name]
 con=[x for x in constraints if x['table_name']==name]
 con_by=[]
 for x in table_cols:
  cs=[c for c in con if x['column_name'] in c['definition']]
  con_by.append(cs)
 for i,col in enumerate(table_cols,1):
  dtype=col['udt_name'] if col['data_type']=='USER-DEFINED' else col['data_type']
  if dtype=='USER-DEFINED': dtype='public.'+dtype
  size=[]
  if col['character_maximum_length']: size.append('max '+str(col['character_maximum_length']))
  if col['numeric_precision']: size.append('precision '+str(col['numeric_precision'])+(' scale '+str(col['numeric_scale']) if col['numeric_scale'] is not None else ''))
  if col['datetime_precision'] and col['data_type'] in ('timestamp with time zone','timestamp without time zone','time without time zone','time with time zone'):
   size.append('fractional seconds '+str(col['datetime_precision']))
  cset=con_by[i-1]
  pk=any(c['kind']=='p' and c['definition'].startswith('PRIMARY KEY') for c in cset)
  fks=[c for c in cset if c['kind']=='f']
  checks=[c for c in cset if c['kind']=='c']
  uniques=[c for c in cset if c['kind']=='u']
  notes=[]
  if col['description']: notes.append(col['description'])
  if fks: notes.extend('FK '+x['name']+': '+x['definition'] for x in fks)
  if checks: notes.extend('CHECK '+x['name']+': '+x['definition'] for x in checks)
  if uniques: notes.extend('UNIQUE '+x['name']+': '+x['definition'] for x in uniques)
  if not notes: notes.append('See Constraints and Relationships sections; no column comment recorded.' if not col['description'] else col['description'])
  values=[i,col['column_name'],dtype,'; '.join(size) or '—','NO' if col['is_nullable']=='NO' else 'YES','YES' if pk else 'NO',col['column_default'] or '—','; '.join(notes)]
  for j,val in enumerate(values,1):
   cell=ws.cell(6+i,j,val);cell.alignment=Alignment(vertical='top',wrap_text=True)
 end=6+len(table_cols); row=end+2
 sections=[('Indexes',[x for x in indexes if x['table_name']==name],['Index name','Definition','']),
 ('Constraints',con,['Constraint name','Kind','Definition']),
 ('Relationships',[x for x in con if x['kind']=='f'],['Constraint name','Local columns','Referenced relation / columns and action']),
 ('RLS policies',[x for x in policies if x['table_name']==name],['Policy','Command / roles','USING','WITH CHECK']),
 ('Triggers',[x for x in triggers if x['table_name']==name],['Trigger','Event / timing','Definition'])]
 for title,records,headers in sections:
  ws.cell(row,1,title).font=Font(bold=True,size=12,color='17365D'); head=row+1
  for j,h in enumerate(headers,1):
   c=ws.cell(head,j,h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='5B9BD5');c.alignment=Alignment(vertical='top',wrap_text=True)
  row=head+1
  for rec in records:
   if title=='Indexes': vals=[rec['name'],rec['definition'],'']
   elif title=='Constraints': vals=[rec['name'],{'p':'PRIMARY KEY','f':'FOREIGN KEY','u':'UNIQUE','c':'CHECK','x':'EXCLUSION'}.get(rec['kind'],rec['kind']),rec['definition']]
   elif title=='Relationships': vals=[rec['name'],rec['definition'].split('REFERENCES')[0].replace('FOREIGN KEY','').strip(),rec['definition'].split('REFERENCES',1)[1] if 'REFERENCES' in rec['definition'] else rec['definition']]
   elif title=='RLS policies': vals=[rec['name'],rec['command']+' / '+', '.join(rec['roles']),rec['using_expression'] or '—',rec['check_expression'] or '—']
   else: vals=[rec['name'],rec['event']+' / '+rec['timing']+' / '+rec['orientation'],rec['definition']]
   for j,v in enumerate(vals,1): ws.cell(row,j,str(v) if v not in (None,'') else '—').alignment=Alignment(vertical='top',wrap_text=True)
   row+=1
  if not records: ws.cell(row,1,'None recorded for this table.').alignment=Alignment(wrap_text=True);row+=1
  row+=1
 status=next(x for x in rls if x['table_name']==name)
 ws.cell(row,1,'Row Level Security').font=Font(bold=True,color='17365D');ws.cell(row,2,('ENABLED' if status['enabled'] else 'DISABLED')+'; forced='+str(status['forced']).upper())
 ws.freeze_panes='A7';ws.sheet_view.showGridLines=False
 ws.auto_filter.ref=f'A6:H{end}'
 for col_letter,width in {'A':11,'B':27,'C':25,'D':25,'E':19,'F':23,'G':38,'H':80}.items(): ws.column_dimensions[col_letter].width=width
 ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True

types=wb.create_sheet('Enums & Types');types.append(['Type name','PostgreSQL type','Labels in applied DB','Source'])
for e in enums: types.append([e['name'],e['pg_type'],', '.join(e['values']),'Applied local database catalog'])
for c in types[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78')
types.freeze_panes='A2';types.auto_filter.ref=f'A1:D{types.max_row}'
for col,width in {'A':29,'B':30,'C':75,'D':55}.items():types.column_dimensions[col].width=width

prov=wb.create_sheet('Migration Provenance');prov.append(['Applied version','Migration name','Ledger statements','Repository source and pinned revision','SHA-256 of source SQL','Applied schema cross-check'])
for c in prov[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(wrap_text=True)
for m in migrations:
 match=next((s for s in migration_sources if s[0].startswith(m['version'])),None)
 assert match, f"unmapped applied migration {m}"
 assert int(m['statement_count'])>0
 prov.append([m['version'],m['name'],m['statement_count'],match[1]+' — '+match[2],match[3],'Applied migration ledger; catalog inventory below'])
for row in [
 ['Scope','Local Supabase PostgreSQL schema only; no application/user rows inspected.'],
 ['Database identity','Docker container '+DB+'; PostgreSQL 15.8; Supabase migration ledger queried read-only.'],
 ['Applied objects','11 public tables; 100 columns; 39 indexes; 67 constraints; 6 RLS policies; 7 triggers; 4 public enum types; 7 documented application functions.'],
 ['Migration boundary','Exactly five migration versions in this local DB. Migration 20261020000142 amenities_seed is not in the applied ledger; seed data is not represented as schema.'],
 ['Repository status','The listing/engagement migrations derive from open PR #7, and the search functions from open PR #6 at the pinned heads above. These source revisions are unmerged; workbook documents the observed local applied schema, not production/main.'],
 ['Excluded objects','Supabase-managed auth.users and platform schemas are not public application tables; extension/system objects excluded except enums/functions used by public schema.'],
 ['Limitation','This isolated local DB is the applied schema baseline inspected; no remote/production database or staging environment was accessed.']]: prov.append(row)
prov.freeze_panes='A2';prov.auto_filter.ref=f'A1:F{len(migrations)+1}'
for col,width in {'A':24,'B':36,'C':20,'D':78,'E':68,'F':65}.items():prov.column_dimensions[col].width=width

obj=wb.create_sheet('Applied Objects');obj.append(['Object type','Table/name','Definition / values'])
for c in obj[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78')
for e in enums:obj.append(['ENUM','public.'+e['name'],', '.join(e['values'])])
for f in functions:obj.append(['FUNCTION','public.'+f['name']+'('+f['arguments']+')','RETURNS '+f['returns']+'; SECURITY '+f['security']+'; '+f['configuration']+'; '+(f['description'] or 'No function comment recorded.')])
obj.freeze_panes='A2';obj.auto_filter.ref=f'A1:C{obj.max_row}'
for col,width in {'A':22,'B':56,'C':125}.items():obj.column_dimensions[col].width=width

# Fill existing ERD placeholder with all FK relationships, avoiding diagrammatic inference.
erd=wb['ERD'];erd['A1']='Applied-schema relationships (PostgreSQL foreign keys)';erd['A1'].font=Font(bold=True,size=15,color='FFFFFF');erd['A1'].fill=PatternFill('solid',fgColor='17365D')
for i,c in enumerate(['Table','Constraint','Definition'],1):
 cell=erd.cell(3,i,c);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78')
for i,x in enumerate([c for c in constraints if c['kind']=='f'],4):
 erd.cell(i,1,'public.'+x['table_name']);erd.cell(i,2,x['name']);erd.cell(i,3,x['definition'])
 for j in range(1,4): erd.cell(i,j).alignment=Alignment(vertical='top',wrap_text=True)
for col,width in {'A':28,'B':42,'C':100}.items():erd.column_dimensions[col].width=width
erd.freeze_panes='A4';erd.auto_filter.ref=f'A3:C{len([c for c in constraints if c["kind"]=="f"])+3}'
wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
wb.save(OUTPUT)
print(json.dumps({'output':str(OUTPUT),'tables':len(tables),'columns':len(cols),'indexes':len(indexes),'constraints':len(constraints),'policies':len(policies),'triggers':len(triggers),'enums':len(enums),'migrations':len(migrations),'functions':len(functions),'sha256':sha(OUTPUT)},indent=2))
