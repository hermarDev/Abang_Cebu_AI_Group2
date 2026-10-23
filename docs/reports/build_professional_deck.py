"""Evidence-led presentation from the verified 16 Oct 2026 Jira sprint snapshot."""
import ast
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).parent
ROOT = HERE.parent.parent
BASE = 'https://abangcebuai.atlassian.net/browse/'
SOURCE = HERE / 'build_sprint_ticket_deck.py'
OUT = HERE / 'AbangCebu_Sprints_1-3_Professional.pptx'
SNAP = HERE / 'sprint-ticket-snapshot.json'
MANIFEST = HERE / 'sprint-ticket-evidence.json'

# Extract only the literal Jira issue data from the historical builder, not its side effects.
tree = ast.parse(SOURCE.read_text())
literals = {target.id: ast.literal_eval(node.value) for node in tree.body if isinstance(node, ast.Assign) for target in node.targets if isinstance(target, ast.Name) and target.id in ('S1','S2','S3')}
S1, S2, S3 = literals['S1'], literals['S2'], literals['S3']
SPRINTS = [
    {'name':'Sprint 1','id':4,'state':'closed','dates':'15 Sep–2 Oct planned · completed 1 Oct','issues':[{'key':k,'summary':t,'status':'Done'} for k,t in S1]},
    {'name':'Sprint 2','id':70,'state':'active','dates':'1–15 Oct planned · still active on 16 Oct','issues':[{'key':k,'summary':t,'status':s} for k,t,s in S2]},
    {'name':'Sprint 3','id':72,'state':'future','dates':'No Jira start/end dates','issues':[{'key':k,'summary':t,'status':'READY'} for k,t in S3]},
]
ALL_KEYS = {i['key'] for sp in SPRINTS for i in sp['issues']}
assert len(ALL_KEYS)==82 and [len(s['issues']) for s in SPRINTS]==[27,20,35]
EVIDENCE = [
    {'key':'SCRUM-64','asset':'docs/assets/wireframes/AbangCebu_Landing_Page_Desktop_Wireframe_Blueprint.png','caption':'Desktop landing-page blueprint · wireframe, not a live application','source':'Local wireframe; Jira lists same-named attachment','bullets':['Map-led rental discovery layout','Search, filters and listing preview are design concepts']},
    {'key':'SCRUM-65','asset':'docs/assets/wireframes/AbangCebu_Login_Page_Desktop_Wireframe_Blueprint.png','caption':'Desktop login-page blueprint · wireframe, not a live application','source':'Local wireframe; Jira attachment linkage not independently checked','bullets':['Login page visual specification','Auth implementation remains separate work']},
    {'key':'SCRUM-104','asset':'docs/assets/flowcharts/auth-registration-pkce.png','caption':'Registration + PKCE verification flow · diagram excerpt','source':'Local PNG; Jira lists matching editable .drawio attachment','bullets':['Defines registration, validation and duplicate checks','Models email verification and PKCE callback']},
    {'key':'SCRUM-116','asset':'docs/assets/flowcharts/map-master-orchestration.png','caption':'Map-first spatial discovery · diagram excerpt','source':'Local flowchart; attachment linkage not independently checked','bullets':['Map initialization and location choice','Spatial query to price pins and results surfaces']},
    {'key':'SCRUM-118','asset':'docs/assets/flowcharts/map-search-spatial-query.png','caption':'Spatial search query · diagram excerpt','source':'Local PNG; Jira lists matching editable .drawio attachment','bullets':['Landmark auto-suggest and viewport query','PostGIS search specification, not implementation']},
    {'key':'SCRUM-125','asset':'docs/assets/flowcharts/landlord-listing-draft-submit.png','caption':'Landlord listing draft/submission · local draft','source':'Local draft; acceptance and Jira publication not verified','bullets':['Draft and submit listing flow','Ticket remains In Development']},
]

# Keep evidence metadata auditable. This is a fixed 16 Oct snapshot, not an automated live refresh.
if __name__ == '__main__':
    SNAP.write_text(json.dumps({'retrieved_at':'2026-10-16','source':'Jira board 1, sprint issue lists reconciled 27/20/35; fixed snapshot','sprints':SPRINTS},indent=2,ensure_ascii=False)+'\n')
    MANIFEST.write_text(json.dumps(EVIDENCE,indent=2,ensure_ascii=False)+'\n')

NAVY=RGBColor(16,45,51); TEAL=RGBColor(8,126,116); PALE=RGBColor(247,248,246)
WHITE=RGBColor(255,255,255); INK=RGBColor(21,48,54); SLATE=RGBColor(85,107,117)
AMBER=RGBColor(178,107,19); LINE=RGBColor(216,229,227)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def rect(s,x,y,w,h,c,rounded=False):
    z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    z.fill.solid(); z.fill.fore_color.rgb=c; z.line.fill.background(); return z

def text(s,value,x,y,w,h,size=18,color=INK,bold=False,link=None):
    z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=z.text_frame; f.clear(); f.word_wrap=True
    f.margin_left=f.margin_right=Inches(.02); f.margin_top=f.margin_bottom=Inches(.01)
    p=f.paragraphs[0]; p.text=value; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    if link: p.runs[0].hyperlink.address=link
    return z

def slide(kicker,title,sub='',dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY if dark else PALE
    col=WHITE if dark else INK
    rect(s,.64,.53,.07,.47,TEAL if not dark else RGBColor(67,210,185))
    text(s,kicker.upper(),.9,.58,11,.28,11,TEAL if not dark else RGBColor(112,235,209),True)
    title_size = 22 if len(title) > 65 else 26 if len(title) > 48 else 31
    text(s,title,.9,1.02,11.8,.88,title_size,col,True)
    if sub: text(s,sub,.9,1.99,11.7,.29,13,RGBColor(189,215,214) if dark else SLATE)
    rect(s,.9,6.94,11.55,.015,LINE if not dark else SLATE)
    text(s,f'ABANG CEBU AI  /  JIRA SNAPSHOT 16 OCT 2026                                          {len(prs.slides):02}',.9,7.03,11.8,.26,9,SLATE if not dark else RGBColor(188,216,214))
    return s

def badge(s,label,x,y,color):
    rect(s,x,y,1.5,.39,color,True); text(s,label,x+.12,y+.07,1.32,.2,10,WHITE,True)

def figure(s,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im: iw,ih=im.size
    # Fit, never distort or crop away chart branches; captions explain full linked source.
    scale=min(w/iw,h/ih); dw,dh=iw*scale,ih*scale
    s.shapes.add_picture(str(path), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))

def divider(n):
    sp=SPRINTS[n]; counts=Counter(i['status'] for i in sp['issues']); name=sp['name']; state=sp['state']
    s=slide(f'{name} / {state}', {'Sprint 1':'Foundation, designed','Sprint 2':'Flows, specified','Sprint 3':'Build, then verify'}[name],sp['dates'],True)
    text(s,str(len(sp['issues'])),1,2.55,3,1.15,62,WHITE,True)
    text(s,'tickets in sprint',1,3.7,4,.4,20,RGBColor(186,214,211))
    x=5.25
    for status,count in counts.items():
        c=RGBColor(81,212,179) if status=='Done' else AMBER if status=='IN DEVELOPMENT' else RGBColor(169,188,195)
        text(s,str(count),x,2.52,2,.9,50,c,True); text(s,status.title(),x,3.44,2.7,.6,17,WHITE,True); x+=2.65
    text(s,'Jira status ≠ independent acceptance or implementation audit.',1,5.49,10,.5,17,RGBColor(186,214,211))

def narrative(kicker,title,items,small=''):
    s=slide(kicker,title,small)
    for i,(label,body) in enumerate(items):
        y=2.55+i*1.31; rect(s,.9,y,11.5,1.11,WHITE,True)
        text(s,label,1.15,y+.14,3.25,.34,18,TEAL,True)
        text(s,body,4.25,y+.13,7.78,.75,17,INK)

def spotlight(entry):
    issue=next(i for sp in SPRINTS for i in sp['issues'] if i['key']==entry['key'])
    key=entry['key']; status=issue['status']; sprint=next(sp['name'] for sp in SPRINTS if issue in sp['issues'])
    s=slide(f'{sprint} / ARTIFACT SPOTLIGHT',key+' · '+issue['summary'], 'Ticket contents and artifact provenance')
    rect(s,.9,2.36,7.0,4.27,WHITE,True)
    figure(s,ROOT/entry['asset'],1.04,2.48,6.72,3.47)
    text(s,entry['caption'],1.1,6.05,6.65,.48,12,SLATE)
    rect(s,8.18,2.36,4.24,4.27,WHITE,True)
    badge(s,'IN DEV' if status=='IN DEVELOPMENT' else status.upper(),8.49,2.64,AMBER if status=='IN DEVELOPMENT' else TEAL)
    y=3.27
    for bullet in entry['bullets']:
        text(s,'• '+bullet,8.49,y,3.57,.68,17,INK); y+=.87
    text(s,entry['source'],8.49,5.25,3.6,.68,11,SLATE)
    text(s,'OPEN JIRA ISSUE ↗',8.49,6.13,3.6,.28,12,TEAL,True,BASE+key)

def inventory(sp):
    issues=sp['issues']; name=sp['name']; total=(len(issues)+5)//6
    for idx,start in enumerate(range(0,len(issues),6),1):
        chunk=issues[start:start+6]
        s=slide(f'INVENTORY / {name}',f'{name} · ticket register',f'{start+1}–{start+len(chunk)} of {len(issues)}  /  appendix page {idx} of {total}')
        for row,issue in enumerate(chunk):
            y=2.33+row*.72
            rect(s,.9,y,11.51,.64,WHITE,True)
            text(s,issue['key'],1.12,y+.13,1.36,.32,14,TEAL,True,BASE+issue['key'])
            text(s,issue['summary'],2.57,y+.1,8.28,.48,14,INK)
            status=issue['status']; short='IN DEV' if status=='IN DEVELOPMENT' else status.upper()
            text(s,short,10.96,y+.17,1.23,.25,10,AMBER if status=='IN DEVELOPMENT' else TEAL if status=='Done' else SLATE,True)

s=slide('Delivery review / 01','From specification to implementation','Sprint-by-sprint Jira tickets · 16 October 2026',True)
text(s,'01',1,3.06,2,.8,54,RGBColor(112,235,209),True); text(s,'Foundation',1,4.01,3,.44,20,WHITE)
text(s,'02',5,3.06,2,.8,54,RGBColor(112,235,209),True); text(s,'Detailed flows',5,4.01,3,.44,20,WHITE)
text(s,'03',9,3.06,2,.8,54,RGBColor(169,188,195),True); text(s,'Planned build',9,4.01,3,.44,20,WHITE)

divider(0)
narrative('Sprint 1 / outcome','Requirements before implementation',[
    ('Foundation','Git, AWS, Next.js/MapLibre and Supabase/PostGIS direction.'),
    ('Access + data','User/profile schema, RLS policy drafts, RBAC and renter/landlord personas.'),
    ('Experience + QA','Page wireframes, authentication workflow specifications and test plans.')],
    '27 of 27 tickets marked Done; most outputs are plans and design artifacts.')
for e in EVIDENCE[:2]: spotlight(e)
inventory(SPRINTS[0])

divider(1)
narrative('Sprint 2 / outcome','Authentication and discovery mapped',[
    ('Authentication','Registration, login, reset, sign-out and failure flows; test matrices.'),
    ('Spatial discovery','Map initialization, search, filters and pin/drawer interaction flowcharts.'),
    ('Carry-over','SCRUM-125, SCRUM-126 and SCRUM-127 remain In Development.')],
    '17 Done; 3 In Development. The sprint still reads Active in Jira.')
for e in EVIDENCE[2:]: spotlight(e)
inventory(SPRINTS[1])

divider(2)
narrative('Sprint 3 / planned work','Implementation has not started in Jira',[
    ('Platform + data','EC2 staging, CI/CD, IAM, monitoring, migrations, RLS and PostGIS.'),
    ('Application','Authentication, RBAC, shared UI, map search and listing interfaces.'),
    ('Verification','Carry-over review, auth/RBAC QA, accessibility, load and system tests.')],
    '35 tickets Ready; 0 Done. No Jira start or end date recorded.')
inventory(SPRINTS[2])

s=slide('Method / source','How to read this deck','Snapshot: 16 Oct 2026 · board 1 · sprint IDs 4, 70, 72')
text(s,'All 82 issue keys appear once in the appendix and link to Jira.',1,2.48,11,.69,23,TEAL,True)
text(s,'The status labels are current Jira workflow states at snapshot time. They are not proof of code, tests, approval, or deployment. Visuals are local artifacts and are captioned with their verification limits.',1,3.35,10.8,1.55,20,INK)
text(s,'Jira board: https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1/backlog',1,5.73,11,.5,13,SLATE)
if __name__ == '__main__':
    prs.save(OUT)
    print(f'{OUT}: {len(prs.slides)} slides, {len(ALL_KEYS)} unique ticket keys, {len(EVIDENCE)} visuals')
