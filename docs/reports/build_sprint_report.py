#!/usr/bin/env python3
"""Generate a source-linked Jira snapshot report; task data transcribed from Jira sprint endpoints."""
from collections import Counter, defaultdict
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
BASE = 'https://abangcebuai.atlassian.net/browse/'
# key | assignee | Jira status | Jira issue summary. All three sprint endpoint pages were checked against totals.
DATA = {
1: '''SCRUM-10|Edrich Bardilas|Done|Create AbangCebu System Flowcharts
SCRUM-49|Hermar centillas|Done|Establish Git Branching & PR Collaboration Strategy
SCRUM-6|Michelle Estoy|Done|create AWS account for the team
SCRUM-55|Michelle Estoy|Done|Draft Row Level Security (RLS) Policy Specifications
SCRUM-62|Michelle Estoy|Done|Define Landlord Persona, Permissions & Capabilities
SCRUM-50|Angel Crushein Yaun|Done|Configure UI Framework & Styling Guidelines
SCRUM-51|Edrich Bardilas|Done|Environment Variables & Secrets Governance Plan
SCRUM-52|Dency Marie Bosque|Done|Plan Supabase Project Structure & Extensions
SCRUM-53|Dency Marie Bosque|Done|Design Initial Database ERD Diagram
SCRUM-54|Dency Marie Bosque|Done|Define Users & Profiles Table Schema Specification
SCRUM-56|John Lloyd Ando|Done|Specify User Registration Workflow & Data Contract
SCRUM-57|John Lloyd Ando|Done|Specify User Login & Session Token Lifecycle
SCRUM-58|junrilldisoy90|Done|Specify User Logout & Session Invalidation
SCRUM-60|junrilldisoy90|Done|Document Authentication Error Handling & Edge Cases
SCRUM-48|Hermar centillas|Done|Plan Next.js 15 & MapLibre Project Foundation
SCRUM-64|Angel Crushein Yaun|Done|Design Landing Page Wireframe / Mockup
SCRUM-63|Joan Marie Encallado Inting|Done|Create Role-Based Access Control (RBAC) Matrix
SCRUM-61|Joan Marie Encallado Inting|Done|Define Renter Persona, Permissions & Capabilities
SCRUM-59|Joan Marie Encallado Inting|Done|Specify Password Reset & Recovery Workflow
SCRUM-65|Neah Moneva|Done|Design Login Page Wireframe / Mockup
SCRUM-66|Angel Crushein Yaun|Done|Design Registration Page Wireframe / Mockup
SCRUM-67|Neah Moneva|Done|Design User Dashboard Wireframe / Mockup
SCRUM-68|Angel Crushein Yaun|Done|Design Profile Page Wireframe / Mockup
SCRUM-69|Jenny Villamor|Done|Draft Authentication Test Plan & Test Scenarios
SCRUM-70|Anne KC M. Casinay|Done|Draft Role & Access Permission Verification Test Cases
SCRUM-71|Karla Hiyas|Done|Draft Protected Route & Middleware Test Scenarios
SCRUM-72|Ryza Albiso|Done|Plan Authentication Baseline Latency & Load Strategy''',
2: '''SCRUM-104|John Lloyd Ando|Done|Sign Up Module & Flowchart: Registration & PKCE Verification
SCRUM-105|Neah Moneva|Done|Log In Module & Flowchart: Authentication & Session Token Lifecycle
SCRUM-106|Joan Marie Encallado Inting|Done|Password Reset Module & Flowchart: Self-Service Recovery Workflow
SCRUM-107|junrilldisoy90|Done|Sign Out Module & Flowchart: Session Termination & Multi-Tab Sync
SCRUM-108|junrilldisoy90|Done|Error Handling & Suspension Module & Flowchart: Failure Taxonomy & UX States
SCRUM-109|Jenny Villamor|Done|Authentication Master Test Plan & Test Matrix
SCRUM-110|Jenny Villamor|Done|Sign Up Module Test Plan & Registration Test Cases
SCRUM-111|Karla Hiyas|Done|Log In Module Test Plan & Session Token Test Cases
SCRUM-112|Anne KC M. Casinay|Done|Password Reset & Sign Out Module Test Cases
SCRUM-113|Karla Hiyas|Done|Route Guards & Middleware Security Test Cases
SCRUM-114|Ryza Albiso|Done|Authentication Baseline Latency & Load Strategy
SCRUM-116|Hermar centillas|Done|Landing Page Master Orchestration & Spatial Discovery Flowchart
SCRUM-117|Hermar centillas|Done|Sub-Process 2.1: Map Initialization, Viewport & Geolocation Flowchart
SCRUM-118|Hermar centillas|Done|Sub-Process 2.2: Spatial Search, Landmark Auto-Suggest & PostGIS Query Flowchart
SCRUM-119|Hermar centillas|Done|Sub-Process 2.3: Multi-Criteria Filter Pills & URL State Sync Flowchart
SCRUM-120|Hermar centillas|Done|Sub-Process 2.4: Bidirectional Pin Marker & 3-Snap Gesture Drawer Flowchart
SCRUM-122|Edrich Bardilas|Done|FUNCTION SPECIFICATION DOCUMENT FOR ABANG CEBU
SCRUM-125|Hermar centillas|IN DEVELOPMENT|Module 3: Landlord Onboarding & Listing Management Flowcharts
SCRUM-126|Hermar centillas|IN DEVELOPMENT|Module 4: Admin KYC Verification & Listing Moderation Flowcharts
SCRUM-127|Hermar centillas|IN DEVELOPMENT|Module 5: Renter Favorites, Contact & Viewing Inquiry Flowcharts''',
3: '''SCRUM-134|Hermar centillas|READY|Provision AWS EC2 staging environment for Next.js app
SCRUM-135|Hermar centillas|READY|Set up GitHub Actions CI pipeline (tsc, lint, build)
SCRUM-136|Hermar centillas|READY|Automate deployment from main to EC2 staging
SCRUM-137|Hermar centillas|READY|Define IAM least-privilege roles and policies for staging
SCRUM-138|Hermar centillas|READY|IAM access review and security audit of staging
SCRUM-139|Hermar centillas|READY|Set up monitoring, logging and uptime alerts for staging
SCRUM-140|Hermar centillas|READY|Implement core listing tables migration
SCRUM-141|Hermar centillas|READY|Implement engagement tables migration (inquiries, saved_listings, reviews, audit_logs)
SCRUM-142|Hermar centillas|READY|Apply Row Level Security policies as migrations
SCRUM-143|Hermar centillas|READY|Create PostGIS search functions (viewport bbox and landmark radius)
SCRUM-144|Hermar centillas|READY|Seed Metro Cebu landmarks and sample listings
SCRUM-145|Hermar centillas|READY|Configure Supabase Storage buckets and database backups
SCRUM-146|Hermar centillas|READY|Complete Database Specification workbook (.xlsx)
SCRUM-147|Hermar centillas|READY|Implement sign-up and email verification (PKCE)
SCRUM-148|Hermar centillas|READY|Implement login, session refresh and Remember Me
SCRUM-149|Hermar centillas|READY|Implement logout and multi-tab sign-out sync
SCRUM-150|Hermar centillas|READY|Implement password reset and recovery
SCRUM-151|Hermar centillas|READY|Implement RBAC middleware and route guards
SCRUM-152|Hermar centillas|READY|Implement standard API error contract and auth error handling
SCRUM-153|Hermar centillas|READY|Build shared UI component library (Tailwind v4 tokens)
SCRUM-154|Hermar centillas|READY|Implement login, registration and password reset pages
SCRUM-155|Hermar centillas|READY|Implement MapLibre map canvas with Metro Cebu viewport and Locate Me
SCRUM-156|Hermar centillas|READY|Implement landing page, search bar and landmark auto-suggest
SCRUM-157|Hermar centillas|READY|Implement filter pills, URL state sync and results drawer
SCRUM-158|Hermar centillas|READY|Module 6: AI Rental Assistant Flowcharts and Specification
SCRUM-159|Hermar centillas|READY|Define AI assistant API contract, prompt policy and evaluation set
SCRUM-160|Hermar centillas|READY|Extend Function Specification Document with Modules 3 to 5
SCRUM-161|Hermar centillas|READY|Review and close Sprint 2 carry-over (SCRUM-125, 126, 127)
SCRUM-162|Hermar centillas|READY|Execute authentication test plan against staging
SCRUM-163|Hermar centillas|READY|Execute RBAC and route guard tests
SCRUM-164|Hermar centillas|READY|Browser, mobile and usability testing of auth pages and map UI
SCRUM-165|Hermar centillas|READY|Write test plans for Modules 3, 4 and 5
SCRUM-166|Hermar centillas|READY|Execute k6 load tests on staging (steady-state and burst)
SCRUM-167|Hermar centillas|READY|System test plan and first end-to-end run (auth to map to listing)
SCRUM-168|Hermar centillas|READY|Sprint 3 review, retrospective and reporting'''
}
issues = {n: [line.split('|', 3) for line in text.splitlines()] for n, text in DATA.items()}
assert [len(issues[n]) for n in (1,2,3)] == [27,20,35]
assert len({row[0] for rows in issues.values() for row in rows}) == 82

D = Document(); sec=D.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
styles=D.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9)
styles['Normal'].paragraph_format.space_after=Pt(5)
D.add_heading('AbangCebu AI | Sprints 1–3', 0)
D.add_paragraph('Jira delivery and assignment report • Snapshot: 16 October 2026 (Philippines time)')
D.add_paragraph('Project: Abang Cebu AI (SCRUM) • Board: SCRUM board #1 • 82 sprint issue entries across three sprints [1–3].')
D.add_heading('Executive summary', 1)
D.add_paragraph('Sprint 1 (15 September–2 October planned) is closed, with all 27 listed issues currently Done. It concentrated on system and technical planning: account and repository setup, Supabase/PostGIS schema and access policy specifications, identity workflows and role definitions, UI wireframes, and test plans. Closing occurred on 1 October. [1]')
D.add_paragraph('Sprint 2 began 1 October and had a planned end of 15 October. Jira still labels it active: 17 of 20 issues are Done; landlord listing, admin moderation, and renter engagement flowchart tickets (SCRUM-125–127) remain IN DEVELOPMENT. Its work centered on authentication/module diagrams, test specifications, spatial-discovery flowcharts, and the functional specification. [2, 4]')
D.add_paragraph('Sprint 3 is marked future, with no Jira start or end date. All 35 tickets are READY and assigned to Hermar centillas; they describe intended infrastructure, database migrations, authentication, map/search UI, AI-assistant planning, and testing—not completed work. Sprint 2 carry-over is explicitly tracked in SCRUM-161. [3]')
D.add_heading('At-a-glance status', 1)
t=D.add_table(rows=1, cols=4); t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Sprint / Jira ID','Lifecycle','Current issue status','Dates (Jira)']): c.text=v
for vals in [('1 / #4','Closed','27 Done / 27','15 Sep–2 Oct 2026; closed 1 Oct'),('2 / #70','Active','17 Done; 3 IN DEVELOPMENT / 20','1–15 Oct 2026; past planned end'),('3 / #72','Future','35 READY / 35','No start/end recorded')]:
 for c,v in zip(t.add_row().cells,vals): c.text=v
D.add_paragraph('Status is a current Jira snapshot, not a historical burndown. A Done ticket is not independent proof of implementation, testing, or acceptance; READY and future do not imply execution. Issue titles are reproduced as Jira records; assignee means assigned owner, not independently verified author. [1–4]')

for n in (1,2,3):
 D.add_page_break(); D.add_heading(f'Sprint {n}: narrative and people', 1)
 if n==1:
  D.add_paragraph('Foundation and requirements sprint. Planning established a Next.js/MapLibre direction and Git collaboration approach; database and Supabase design, RLS, personas and RBAC defined the intended data/security model. Identity workflow specifications and page wireframes fed later module work, while testing and load-strategy tickets laid out validation. All 27 current ticket states are Done, but this report does not claim those plans were executed in production. [1]')
 elif n==2:
  D.add_paragraph('Detailed design and validation planning sprint. Auth sign-up/login/reset/logout and error-state diagrams were paired with test matrices. Map initialization, search, filter and drawer flows were drawn up. The functional specification was marked Done. Module 3–5 flowchart work remains IN DEVELOPMENT; local audit notes say draft packages were attached, not accepted, and visual/product/schema review remained open. [2, 4]')
 else:
  D.add_paragraph('Prospective implementation and QA backlog, not delivery history. The 35 READY tickets encompass EC2/CI/IAM/monitoring, database and RLS migrations, auth and RBAC implementation, page/map/search development, AI assistant design, and staging/security/performance/system tests. All are assigned to one person in Jira; this is an assignment concentration, not evidence that one person will perform all work. [3]')
 counts=Counter(r[1] for r in issues[n]); D.add_heading('Assigned people', 2)
 t=D.add_table(rows=1, cols=2); t.style='Light Shading Accent 1'; t.rows[0].cells[0].text='Jira assignee'; t.rows[0].cells[1].text='Tickets'
 for who,num in sorted(counts.items(), key=lambda x:(-x[1],x[0])):
  cells=t.add_row().cells; cells[0].text=who; cells[1].text=str(num)
 D.add_heading('Complete issue register', 2)
 D.add_paragraph('Each key resolves at https://abangcebuai.atlassian.net/browse/<key>. Titles, owners and states were read from the Jira sprint-issues endpoint. [1–3]')
 t=D.add_table(rows=1, cols=3); t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,['Issue and task','Assignee','State']): c.text=v
 for key,person,status,title in issues[n]:
  cells=t.add_row().cells; cells[0].text=f'{key} — {title}'; cells[1].text=person; cells[2].text=status
  for c in cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1)
    for run in p.runs: run.font.size=Pt(7.5)

D.add_page_break(); D.add_heading('Cross-sprint interpretation and risks', 1)
for text in [
 'Continuity: Sprint 1 specified identity, data and access control; Sprint 2 expanded those into flowcharts and test plans; Sprint 3 proposes implementing and validating those foundations. This is a thematic reading of the issue titles, not a verified dependency graph. [1–3]',
 'Outstanding work: Sprint 2 is still active beyond its planned end, with SCRUM-125, 126 and 127 IN DEVELOPMENT. An audit dated 16 October records draft ZIP attachments but says acceptance and visual/product/schema QA remain open. The Sprint 3 carry-over ticket SCRUM-161 acknowledges those three. [2, 3, 4]',
 'Capacity: all 35 Sprint 3 issues are assigned to Hermar centillas. Before starting, revalidate estimates, ownership, capacity and the sequence of staging, migrations, auth, UI and testing. No velocity or actual effort was verified in this report. [3]',
 'Evidence boundary: This is a Jira issue/status/assignee snapshot plus a local audit note; it is not a review of each attachment, code commit, acceptance criterion, or historical status transition. “Done” is reported as Jira says, not certified delivery. [1–4]']:
 D.add_paragraph(text, style='List Bullet')
D.add_heading('Sources and method', 1)
for s in [
 '[1] Jira SCRUM board #1, Sprint 1 #4, sprint metadata and all 27 sprint issues (retrieved 16 October 2026): https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1',
 '[2] Jira SCRUM board #1, Sprint 2 #70, sprint metadata and all 20 sprint issues (retrieved 16 October 2026): https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1',
 '[3] Jira SCRUM board #1, Sprint 3 #72, sprint metadata and all 35 sprint issues (retrieved 16 October 2026): https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1',
 '[4] Local project audit, “AbangCebu AI Sprint 2 Coverage and Sprint 3 Planning Report,” dated 16 October 2026: docs/specifications/sprint2-final-audit.md (internal working document; secondary evidence).',
 'Individual issue source: https://abangcebuai.atlassian.net/browse/SCRUM-<number>. Counts reconciled to sprint endpoint totals (27 + 20 + 35 = 82); assignments and current statuses were counted from all entries.']:
 D.add_paragraph(s)
footer=sec.footer.paragraphs[0]; footer.text='AbangCebu AI • Jira snapshot • 16 October 2026'
out=ROOT/'AbangCebu_Sprints_1-3_Report.docx'; D.save(out); print(out)
