"""Build Jira sprint ticket presentation from the verified 16 Oct 2026 snapshot."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).with_name('AbangCebu_Sprints_1-3_Ticket_Status.pptx')
BASE = 'https://abangcebuai.atlassian.net/browse/'
# Exact Jira sprint issue-list titles and current statuses, retrieved 16 Oct 2026.
S1 = [
('SCRUM-10','Create AbangCebu System Flowcharts'),
('SCRUM-49','Establish Git Branching & PR Collaboration Strategy'),
('SCRUM-6','create AWS account for the team'),
('SCRUM-55','Draft Row Level Security (RLS) Policy Specifications'),
('SCRUM-62','Define Landlord Persona, Permissions & Capabilities'),
('SCRUM-50','Configure UI Framework & Styling Guidelines'),
('SCRUM-51','Environment Variables & Secrets Governance Plan'),
('SCRUM-52','Plan Supabase Project Structure & Extensions'),
('SCRUM-53','Design Initial Database ERD Diagram'),
('SCRUM-54','Define Users & Profiles Table Schema Specification'),
('SCRUM-56','Specify User Registration Workflow & Data Contract'),
('SCRUM-57','Specify User Login & Session Token Lifecycle'),
('SCRUM-58','Specify User Logout & Session Invalidation'),
('SCRUM-60','Document Authentication Error Handling & Edge Cases'),
('SCRUM-48','Plan Next.js 15 & MapLibre Project Foundation'),
('SCRUM-64','Design Landing Page Wireframe / Mockup'),
('SCRUM-63','Create Role-Based Access Control (RBAC) Matrix'),
('SCRUM-61','Define Renter Persona, Permissions & Capabilities'),
('SCRUM-59','Specify Password Reset & Recovery Workflow'),
('SCRUM-65','Design Login Page Wireframe / Mockup'),
('SCRUM-66','Design Registration Page Wireframe / Mockup'),
('SCRUM-67','Design User Dashboard Wireframe / Mockup'),
('SCRUM-68','Design Profile Page Wireframe / Mockup'),
('SCRUM-69','Draft Authentication Test Plan & Test Scenarios'),
('SCRUM-70','Draft Role & Access Permission Verification Test Cases'),
('SCRUM-71','Draft Protected Route & Middleware Test Scenarios'),
('SCRUM-72','Plan Authentication Baseline Latency & Load Strategy'),
]
S2 = [
('SCRUM-104','Sign Up Module & Flowchart: Registration & PKCE Verification','Done'),
('SCRUM-105','Log In Module & Flowchart: Authentication & Session Token Lifecycle','Done'),
('SCRUM-106','Password Reset Module & Flowchart: Self-Service Recovery Workflow','Done'),
('SCRUM-107','Sign Out Module & Flowchart: Session Termination & Multi-Tab Sync','Done'),
('SCRUM-108','Error Handling & Suspension Module & Flowchart: Failure Taxonomy & UX States','Done'),
('SCRUM-109','Authentication Master Test Plan & Test Matrix','Done'),
('SCRUM-110','Sign Up Module Test Plan & Registration Test Cases','Done'),
('SCRUM-111','Log In Module Test Plan & Session Token Test Cases','Done'),
('SCRUM-112','Password Reset & Sign Out Module Test Cases','Done'),
('SCRUM-113','Route Guards & Middleware Security Test Cases','Done'),
('SCRUM-114','Authentication Baseline Latency & Load Strategy','Done'),
('SCRUM-116','Landing Page Master Orchestration & Spatial Discovery Flowchart','Done'),
('SCRUM-117','Sub-Process 2.1: Map Initialization, Viewport & Geolocation Flowchart','Done'),
('SCRUM-118','Sub-Process 2.2: Spatial Search, Landmark Auto-Suggest & PostGIS Query Flowchart','Done'),
('SCRUM-119','Sub-Process 2.3: Multi-Criteria Filter Pills & URL State Sync Flowchart','Done'),
('SCRUM-120','Sub-Process 2.4: Bidirectional Pin Marker & 3-Snap Gesture Drawer Flowchart','Done'),
('SCRUM-122','FUNCTION SPECIFICATION DOCUMENT FOR ABANG CEBU','Done'),
('SCRUM-125','Module 3: Landlord Onboarding & Listing Management Flowcharts','IN DEVELOPMENT'),
('SCRUM-126','Module 4: Admin KYC Verification & Listing Moderation Flowcharts','IN DEVELOPMENT'),
('SCRUM-127','Module 5: Renter Favorites, Contact & Viewing Inquiry Flowcharts','IN DEVELOPMENT'),
]
S3 = [
('SCRUM-134','Provision AWS EC2 staging environment for Next.js app'),
('SCRUM-135','Set up GitHub Actions CI pipeline (tsc, lint, build)'),
('SCRUM-136','Automate deployment from main to EC2 staging'),
('SCRUM-137','Define IAM least-privilege roles and policies for staging'),
('SCRUM-138','IAM access review and security audit of staging'),
('SCRUM-139','Set up monitoring, logging and uptime alerts for staging'),
('SCRUM-140','Implement core listing tables migration'),
('SCRUM-141','Implement engagement tables migration (inquiries, saved_listings, reviews, audit_logs)'),
('SCRUM-142','Apply Row Level Security policies as migrations'),
('SCRUM-143','Create PostGIS search functions (viewport bbox and landmark radius)'),
('SCRUM-144','Seed Metro Cebu landmarks and sample listings'),
('SCRUM-145','Configure Supabase Storage buckets and database backups'),
('SCRUM-146','Complete Database Specification workbook (.xlsx)'),
('SCRUM-147','Implement sign-up and email verification (PKCE)'),
('SCRUM-148','Implement login, session refresh and Remember Me'),
('SCRUM-149','Implement logout and multi-tab sign-out sync'),
('SCRUM-150','Implement password reset and recovery'),
('SCRUM-151','Implement RBAC middleware and route guards'),
('SCRUM-152','Implement standard API error contract and auth error handling'),
('SCRUM-153','Build shared UI component library (Tailwind v4 tokens)'),
('SCRUM-154','Implement login, registration and password reset pages'),
('SCRUM-155','Implement MapLibre map canvas with Metro Cebu viewport and Locate Me'),
('SCRUM-156','Implement landing page, search bar and landmark auto-suggest'),
('SCRUM-157','Implement filter pills, URL state sync and results drawer'),
('SCRUM-158','Module 6: AI Rental Assistant Flowcharts and Specification'),
('SCRUM-159','Define AI assistant API contract, prompt policy and evaluation set'),
('SCRUM-160','Extend Function Specification Document with Modules 3 to 5'),
('SCRUM-161','Review and close Sprint 2 carry-over (SCRUM-125, 126, 127)'),
('SCRUM-162','Execute authentication test plan against staging'),
('SCRUM-163','Execute RBAC and route guard tests'),
('SCRUM-164','Browser, mobile and usability testing of auth pages and map UI'),
('SCRUM-165','Write test plans for Modules 3, 4 and 5'),
('SCRUM-166','Execute k6 load tests on staging (steady-state and burst)'),
('SCRUM-167','System test plan and first end-to-end run (auth to map to listing)'),
('SCRUM-168','Sprint 3 review, retrospective and reporting'),
]
assert len(S1) == 27 and len(S2) == 20 and len(S3) == 35
assert len({k for k,*_ in S1+S2+S3}) == 82
assert sum(s=='Done' for _,_,s in S2) == 17
NAVY=RGBColor(9,25,41); PANEL=RGBColor(17,39,56); WHITE=RGBColor(240,248,248)
MUTED=RGBColor(159,183,191); TEAL=RGBColor(61,214,177); AMBER=RGBColor(250,190,94)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def box(slide,x,y,w,h,color):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background(); return sh

def txt(slide,text,x,y,w,h,size=18,color=WHITE,bold=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame
    tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=Inches(.03); tf.margin_top=Inches(.02)
    p=tf.paragraphs[0]; p.text=text; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    return sh

def frame(section,title,subtitle=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
    box(s,.48,.48,.08,.62,TEAL); txt(s,section.upper(),.74,.48,11,.3,11,TEAL,True)
    txt(s,title,.74,.86,12, .6,27,WHITE,True)
    if subtitle: txt(s,subtitle,.74,1.49,12,.36,12,MUTED)
    txt(s,'ABANG CEBU AI  •  JIRA SNAPSHOT 16 OCT 2026  •  '+str(len(prs.slides)),.74,7.12,11,.25,9,MUTED)
    return s

def link_key(slide,key,x,y):
    sh=txt(slide,key,x,y,1.2,.25,12,TEAL,True)
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address=BASE+key

def ticket_page(sprint,items,start,end,page,desc):
    s=frame(f'{sprint} / TICKETS {page}',f'{sprint}: tickets {start}–{end}',desc)
    # Two columns, five cards per column. Cards show exact Jira title and status.
    for i,item in enumerate(items):
        key,title,status=item
        col=i//5; row=i%5; x=.74+col*6.13; y=1.98+row*.98
        box(s,x,y,5.87,.89,PANEL)
        link_key(s,key,x+.14,y+.1)
        label='IN DEV' if status=='IN DEVELOPMENT' else status.upper()
        c=AMBER if status=='IN DEVELOPMENT' else TEAL if status=='Done' else MUTED
        txt(s,label,x+4.67,y+.12,1.04,.24,9,c,True)
        txt(s,title,x+.14,y+.39,5.55,.43,12,WHITE)
    return s

def pages(sprint,items,desc):
    for start in range(0,len(items),10):
        ticket_page(sprint,items[start:start+10],start+1,min(start+10,len(items)),start//10+1,desc)

s=frame('Sprint 1 → Sprint 2 → Sprint 3','What is Done? What is Next?','Abang Cebu AI • ticket-by-ticket Jira status • 16 October 2026')
txt(s,'27 / 27',.9,2.55,3.7,1,40,TEAL,True); txt(s,'SPRINT 1  •  DONE',.9,3.52,3.8,.35,16,WHITE,True)
txt(s,'17 / 20',5,2.55,3.7,1,40,TEAL,True); txt(s,'SPRINT 2  •  DONE',5,3.52,3.8,.35,16,WHITE,True)
txt(s,'0 / 35',9.05,2.55,3.6,1,40,MUTED,True); txt(s,'SPRINT 3  •  DONE',9.05,3.52,3.6,.35,16,WHITE,True)
txt(s,'Jira status is not independent proof of implementation or acceptance.',.9,5.78,11,.45,15,MUTED)

s=frame('Sprint 1 / closed','27 of 27 tickets Done','15 Sep–2 Oct planned • completed 1 Oct 2026 • foundation, requirements, UX, test plans')
txt(s,'All 27 Sprint 1 issue statuses are Done.',.9,2.32,11,.8,30,TEAL,True)
txt(s,'Scope: project setup; database and security design; authentication specifications; wireframes; test and performance planning.',.9,3.55,11,1.5,22,WHITE)
pages('Sprint 1',[(k,t,'Done') for k,t in S1],'27 / 27 Done • Jira sprint closed')

s=frame('Sprint 2 / active','17 Done · 3 In Development','1–15 Oct planned • Jira still marks the sprint active as of 16 Oct 2026')
txt(s,'Done: authentication flows and test plans, spatial discovery flows, functional specification.',.9,2.22,11,1.15,22,TEAL,True)
txt(s,'Not Done: SCRUM-125 landlord; SCRUM-126 admin KYC; SCRUM-127 renter engagement flowcharts.',.9,3.65,11,1.4,22,AMBER,True)
pages('Sprint 2',S2,'17 Done • 3 In Development • Jira sprint active')

s=frame('Sprint 3 / future','35 Ready · 0 Done','No start or end date recorded in Jira • implementation and verification backlog')
txt(s,'Planned: staging and CI/CD; database/RLS; authentication, map and UI; AI design; security and QA.',.9,2.27,11,1.4,23,WHITE,True)
txt(s,'Ready is a workflow status, not evidence of completed work.',.9,4.14,11,.7,20,AMBER)
pages('Sprint 3',[(k,t,'READY') for k,t in S3],'35 Ready • 0 Done • sprint not started')

s=frame('Source and interpretation','Jira status snapshot','Board 1 • project SCRUM • sprint IDs 4, 70, 72')
txt(s,'82 sprint issue entries: 27 + 20 + 35. Counts reconciled with Jira sprint issue-list totals.',.9,2.2,11,1,23,WHITE,True)
txt(s,'Issue keys on the ticket slides link to Jira. “Done” means the current Jira workflow status only; it is not a code, artifact, test, or acceptance audit.',.9,3.42,11,1.45,19,MUTED)
txt(s,'Source: https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1/backlog',.9,5.36,11,.56,13,TEAL)
OUT.parent.mkdir(parents=True,exist_ok=True); prs.save(OUT)
print(f'{OUT}: {len(prs.slides)} slides; {len(S1)+len(S2)+len(S3)} tickets')
