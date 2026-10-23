"""Build the AbangCebu sprint briefing from the verified Jira snapshot."""
import os
import sys
sys.path.insert(0, '/home/hrmr/.hermes/cache/scratch/pptx-deps')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(__file__), 'AbangCebu_Sprints_1-3_Presentation.pptx')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG = RGBColor(10, 19, 37)
PANEL = RGBColor(21, 35, 58)
WHITE = RGBColor(241, 247, 252)
MUTED = RGBColor(159, 178, 198)
CYAN = RGBColor(65, 211, 215)
GREEN = RGBColor(86, 211, 163)
AMBER = RGBColor(255, 190, 91)
LINE = RGBColor(42, 62, 83)
BOARD = 'https://abangcebuai.atlassian.net/jira/software/projects/SCRUM/boards/1'

def rect(s, x,y,w,h, color, radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color
    sh.line.fill.background()
    if radius:
        sh.adjustments[0] = .14
    return sh

def text(s, value, x,y,w,h, size=20, color=WHITE, bold=False, align=None):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=Inches(.01)
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
        p.space_after=Pt(7)
        if align is not None: p.alignment=align
    return sh

def slide(kicker, title, n, note='Jira snapshot • 16 Oct 2026'):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    rect(s,.52,.42,.07,.34,CYAN)
    text(s,kicker.upper(),.73,.39,11,.4,13,CYAN,True)
    text(s,title,.7,.92,12, .65,32,WHITE,True)
    rect(s,.7,6.99,11.95,.014,LINE)
    text(s,note,.73,7.08,10,.24,10,MUTED)
    text(s,f'{n:02d}',12.18,7.06,.42,.24,10,CYAN,True,PP_ALIGN.RIGHT)
    return s

def card(s,x,y,w,h,heading,body,accent=CYAN):
    rect(s,x,y,w,h,PANEL,True); rect(s,x+.2,y+.24,.055,.35,accent)
    text(s,heading,x+.38,y+.2,w-.55,.47,19,WHITE,True)
    text(s,body,x+.27,y+.83,w-.5,h-.99,16,MUTED)

def bullet(s, y, lead, detail, accent=CYAN):
    rect(s,.82,y+.14,.075,.075,accent)
    text(s,lead,1.06,y,11.3,.37,20,WHITE,True)
    text(s,detail,1.06,y+.41,11.25,.57,15,MUTED)

s=slide('AbangCebu AI / delivery review','Sprints 1–3',1,'Jira issue/status snapshot • 16 Oct 2026 • Philippines time')
text(s,'From foundation and design to the implementation backlog',.78,2.04,10.9,1.02,26,MUTED)
for i,(num,label,c) in enumerate([('27/27','Sprint 1 • Done',GREEN),('17/20','Sprint 2 • Done',AMBER),('0/35','Sprint 3 • Done',CYAN)]):
    x=.78+i*4.12; rect(s,x,3.42,3.82,1.72,PANEL,True)
    text(s,num,x+.26,3.7,3.2,.63,36,c,True)
    text(s,label,x+.26,4.49,3.2,.34,17,WHITE)
text(s,'Status reflects Jira now, not a historical burndown or independent acceptance audit.',.8,5.77,11.5,.62,16,MUTED)

s=slide('01 / executive view','Three sprints, three different states',2)
for i,(head,num,sub,accent) in enumerate([
    ('SPRINT 1  /  CLOSED','27 Done / 27','15 Sep–2 Oct planned; closed 1 Oct',GREEN),
    ('SPRINT 2  /  ACTIVE','17 Done / 20','1–15 Oct planned; 3 in development',AMBER),
    ('SPRINT 3  /  FUTURE','35 Ready / 35','No Jira start or end date recorded',CYAN)]):
    x=.75+i*4.19; rect(s,x,1.93,3.88,3.65,PANEL,True)
    text(s,head,x+.27,2.25,3.38,.44,15,accent,True)
    text(s,num,x+.27,3.05,3.35,.62,27,WHITE,True)
    text(s,sub,x+.27,4.03,3.27,1.07,17,MUTED)
text(s,'82 sprint issue entries total • all three sprint issue lists reconciled to Jira totals.',.78,6.08,11.9,.43,17,WHITE)

s=slide('02 / sprint 1','Foundation and requirements',3)
card(s,.77,1.95,5.75,3.52,'27 of 27 tickets Done','Project setup, Git collaboration, Next.js/MapLibre direction, Supabase/PostGIS schema and RLS planning.',GREEN)
card(s,6.73,1.95,5.78,3.52,'Design handoff','Personas, roles, identity flows, UI wireframes, test plans and performance strategy prepared for subsequent work.',CYAN)
text(s,'Jira state: CLOSED  •  Completed 1 Oct 2026  •  Planned end 2 Oct 2026',.8,5.91,11.6,.54,17,MUTED)

s=slide('03 / sprint 2','Detailed flows and validation plans',4)
card(s,.77,1.92,5.75,3.75,'17 of 20 tickets Done','Authentication diagrams and test matrices; landing/map initialization, spatial search, filters and drawer flows; functional specification.',GREEN)
card(s,6.73,1.92,5.78,3.75,'3 tickets in development','SCRUM-125  Landlord listing flows\nSCRUM-126  Admin KYC / moderation\nSCRUM-127  Renter engagement flows',AMBER)
text(s,'Jira still marks Sprint 2 ACTIVE beyond its planned 15 Oct end.',.8,6.04,11.5,.46,17,AMBER,True)

s=slide('04 / sprint 3','Implementation and QA backlog',5)
card(s,.77,1.93,5.75,3.66,'35 of 35 tickets READY','Planned EC2 staging, CI/CD, IAM, monitoring, database/RLS migrations and authentication/RBAC implementation.',CYAN)
card(s,6.73,1.93,5.78,3.66,'Scope ahead','Planned page, map and search UI; AI-assistant design; security, performance and end-to-end testing.',CYAN)
text(s,'Future sprint; no dates recorded. All 35 tickets assigned to Hermar centillas in Jira.',.8,5.96,11.65,.62,17,MUTED)

s=slide('05 / continuity','How the work connects',6)
for i,(num,label,desc,col) in enumerate([
    ('01','SPECIFY','Identity, data, access, UX',GREEN),
    ('02','DETAIL','Module flows, test design',AMBER),
    ('03','BUILD + VERIFY','Implementation, staging, QA',CYAN)]):
    x=.79+i*4.18; rect(s,x,2.14,3.82,3.13,PANEL,True)
    text(s,num,x+.27,2.38,3.25,.56,27,col,True)
    text(s,label,x+.27,3.22,3.23,.46,21,WHITE,True)
    text(s,desc,x+.27,4.11,3.21,.85,17,MUTED)
text(s,'Thematic reading of Jira titles—not a verified dependency graph.',.81,5.88,11.5,.52,16,MUTED)

s=slide('06 / attention items','Risks to resolve before Sprint 3',7)
bullet(s,1.9,'Close Sprint 2 carry-over','SCRUM-125–127 are still IN DEVELOPMENT; SCRUM-161 tracks the handoff.',AMBER)
bullet(s,3.06,'Confirm acceptance','A Done status is not independent evidence of code, testing or stakeholder acceptance.',CYAN)
bullet(s,4.22,'Rebalance capacity','35 Sprint 3 tickets currently have one Jira assignee; validate owners, estimates and sequence.',AMBER)
text(s,'Local audit note: Sprint 2 draft packages were attached; visual/product/schema QA and acceptance remained open.',1.05,5.62,11.19,.92,15,MUTED)

s=slide('07 / decisions','Recommended next actions',8)
bullet(s,1.93,'1  Review and accept carry-over','Verify the three flowchart tickets against the agreed product/schema criteria.',AMBER)
bullet(s,3.16,'2  Replan Sprint 3','Confirm dates, owners and capacity before starting implementation.',CYAN)
bullet(s,4.39,'3  Set a verification gate','Require ticket-level code, test and acceptance evidence before calling work delivered.',GREEN)

s=slide('08 / sources','Evidence and interpretation',9,'Source retrieval: Jira board #1 • 16 Oct 2026')
text(s,'Jira SCRUM board #1: Sprint 1 (#4), Sprint 2 (#70), Sprint 3 (#72).\n27 + 20 + 35 sprint issue entries; metadata and issue statuses checked live.',.83,2.0,11.6,1.33,20,WHITE)
text(s,BOARD,.83,3.62,11.8,.42,17,CYAN)
text(s,'Supporting narrative: docs/reports/AbangCebu_Sprints_1-3_Report.docx\nScope: issue/status snapshot, not verified burndown, implementation, test results or acceptance.',.83,4.53,11.7,1.16,17,MUTED)
prs.save(OUT)
print(f'{OUT} | {len(prs.slides)} slides | {os.path.getsize(OUT)} bytes')
