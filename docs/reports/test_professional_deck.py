import unittest
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).parent

class ProfessionalDeckTests(unittest.TestCase):
    def test_deck_inventory_and_order(self):
        import build_professional_deck as b
        p = Presentation(ROOT / 'AbangCebu_Sprints_1-3_Professional.pptx')
        texts = ['\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame) for slide in p.slides]
        self.assertGreaterEqual(len(p.slides), 20)
        self.assertEqual(sum('INVENTORY / SPRINT 1' in t for t in texts), 5)
        self.assertEqual(sum('INVENTORY / SPRINT 2' in t for t in texts), 4)
        self.assertEqual(sum('INVENTORY / SPRINT 3' in t for t in texts), 6)
        self.assertLess(next(i for i,t in enumerate(texts) if 'SPRINT 1 / CLOSED' in t), next(i for i,t in enumerate(texts) if 'SPRINT 2 / ACTIVE' in t))
        self.assertLess(next(i for i,t in enumerate(texts) if 'SPRINT 2 / ACTIVE' in t), next(i for i,t in enumerate(texts) if 'SPRINT 3 / FUTURE' in t))
        inv_slides = [s for s,t in zip(p.slides,texts) if 'INVENTORY / SPRINT' in t]
        keys = [r.text for s in inv_slides for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.hyperlink.address and r.text.startswith('SCRUM-')]
        for key in b.ALL_KEYS:
            self.assertEqual(keys.count(key), 1, key)
        self.assertEqual(len(b.ALL_KEYS), 82)

    def test_evidence_and_links(self):
        import build_professional_deck as b
        p = Presentation(ROOT / 'AbangCebu_Sprints_1-3_Professional.pptx')
        for item in b.EVIDENCE:
            self.assertTrue((ROOT.parent.parent / item['asset']).is_file())
            self.assertIn(item['key'], b.ALL_KEYS)
        links = {r.hyperlink.address for s in p.slides for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.hyperlink.address}
        for key in b.ALL_KEYS:
            self.assertIn(b.BASE + key, links)
        self.assertFalse(any('SPRINT 3 / FUTURE' in sh.text and '35 Done' in sh.text for s in p.slides for sh in s.shapes if sh.has_text_frame))

if __name__ == '__main__': unittest.main()
