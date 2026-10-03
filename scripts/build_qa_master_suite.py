import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import markdown
import weasyprint

print("Authoring 100-scenario QA Master Test Suite...")

# 1. 25 Registration Test Cases (SCRUM-110)
reg_cases = [
    {
        "id": "TC-REG-01",
        "title": "Registration - Valid Renter",
        "steps": "1. Submit RegisterPayload with valid email, NIST-compliant password, role='renter', and phone '09171234567'.\n2. Call registration Route Handler.",
        "expected": "1. HTTP 200/201 response.\n2. Phone normalized to canonical +639171234567.\n3. Profile row created in public.profiles via handle_new_user() trigger.\n4. Verification email dispatched.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-02",
        "title": "Registration - Valid Landlord",
        "steps": "1. Submit RegisterPayload with role='landlord', valid credentials, and preferred landmark.\n2. Execute registration Route Handler.",
        "expected": "1. User created with role='landlord'.\n2. Response directs to /landlord/onboarding.\n3. Landlord onboarding profile record initialized.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-03",
        "title": "Registration - Anti-Privilege Escalation",
        "steps": "1. Submit payload attempting role='admin'.\n2. Validate TypeScript contract, Route Handler validation, and database trigger.",
        "expected": "1. Contract rejects payload with INVALID_ROLE.\n2. If bypassed via raw HTTP, server returns HTTP 403 PRIVILEGE_ESCALATION_ATTEMPT.\n3. DB trigger enforces fallback to 'renter'.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-04",
        "title": "Registration - Weak Password Rejection",
        "steps": "1. Attempt registration with password < 8 chars ('Ab1!').\n2. Attempt password missing special character ('Password123').\n3. Attempt password exceeding 72 chars.",
        "expected": "1. Request rejected with HTTP 422 WEAK_PASSWORD.\n2. Response details specifies exact unmet NIST criteria.\n3. Zero user records created.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-05",
        "title": "Registration - Invalid Philippine Phone Rejection",
        "steps": "1. Attempt registration with non-PH numbers ('+1234567890', '08123456789', '12345').",
        "expected": "1. Validation fails with HTTP 422 INVALID_PHONE_NUMBER.\n2. Field error explicitly states requirement for 09XXXXXXXXX format.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-06",
        "title": "Registration - Duplicate Email Handling",
        "steps": "1. Submit registration payload using an existing email address ('mika.cebu@example.com').",
        "expected": "1. Anti-enumeration defense preserves privacy.\n2. Dispatches password reset notification to the owner without revealing account collision to scraper.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-07",
        "title": "Registration - Turnstile Bot Verification Failure",
        "steps": "1. Submit registration request with missing or invalid turnstileToken.",
        "expected": "1. Request immediately rejected with HTTP 400 CAPTCHA_VERIFICATION_FAILED.\n2. Zero database queries executed.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-08",
        "title": "Registration - Duplicate Phone Collision",
        "steps": "1. Submit registration payload using an already registered Philippine phone number with a new email.",
        "expected": "1. Validation catches phone uniqueness conflict.\n2. Returns HTTP 409 PHONE_ALREADY_REGISTERED with user-safe message.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-09",
        "title": "Registration - SQL & Script Injection Sanitization",
        "steps": "1. Enter full_name containing XSS payload: '<script>alert(1)</script>' or SQL fragment \"' OR 1=1 --\".",
        "expected": "1. Payload is fully parameterized and sanitized.\n2. Stored string has tags stripped or HTML escaped.\n3. Zero script execution.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-10",
        "title": "Registration - Email Whitespace & Case Normalization",
        "steps": "1. Submit email with leading/trailing spaces and mixed case: '  Mika.Cebu@Example.COM  '.",
        "expected": "1. Email normalized to lowercase trimmed 'mika.cebu@example.com'.\n2. Stored in canonical form in auth.users.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-REG-11",
        "title": "Registration - Password Confirmation Mismatch",
        "steps": "1. Enter valid password in primary field and differing string in confirm_password field.",
        "expected": "1. Client validator blocks submission with inline error.\n2. If submitted via API, server returns HTTP 422 PASSWORD_CONFIRMATION_MISMATCH.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-12",
        "title": "Registration - Missing Mandatory Fields",
        "steps": "1. Submit JSON payload omitting full_name or role field.",
        "expected": "1. Server returns HTTP 422 VALIDATION_ERROR.\n2. details array contains field-specific path and failure code.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-13",
        "title": "Registration - Rapid Multi-Click Debounce",
        "steps": "1. Rapidly trigger registration submit button 5 times within 300ms.",
        "expected": "1. Client disables submit button immediately after first click.\n2. API receives only single registration dispatch.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-REG-14",
        "title": "Registration - Atomic Transaction Rollback on Failure",
        "steps": "1. Simulate database failure during profile creation trigger after auth.users insertion.",
        "expected": "1. Database transaction rolls back atomically.\n2. Zero orphaned records left in auth.users or public.profiles.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-15",
        "title": "Registration - PKCE Activation Email Dispatch",
        "steps": "1. Complete valid registration.\n2. Inspect outbound transactional email dispatch payload.",
        "expected": "1. Email queued with single-use PKCE verification token.\n2. Redirect URL targets canonical /auth/callback with type=signup.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-16",
        "title": "Verification - Valid PKCE Code Exchange",
        "steps": "1. User clicks confirmation link GET /auth/callback?code=VALID_CODE&type=signup.\n2. Next.js Route Handler exchanges code via @supabase/ssr.",
        "expected": "1. auth.users.email_confirmed_at timestamp set.\n2. Authenticated session cookies issued (Set-Cookie).\n3. Redirects safely to /search (Renter) or /landlord/onboarding (Landlord).",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-17",
        "title": "Verification - Expired Confirmation Code",
        "steps": "1. User clicks confirmation link after 24 hours (code=EXPIRED_CODE).",
        "expected": "1. Route Handler catches Supabase exchange error.\n2. Redirects to /login?error=confirmation_expired.\n3. UI displays persistent banner with resend CTA.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-18",
        "title": "Verification - Tampered / Corrupted PKCE Code",
        "steps": "1. Submit request with corrupted code code=INVALID_12345.",
        "expected": "1. Exchange fails.\n2. Redirects to /login?error=invalid_confirmation_link.\n3. Zero session cookies issued.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-19",
        "title": "Verification - Single-Use Code Replay Invalidation",
        "steps": "1. Click confirmation link once (succeeds).\n2. Attempt to exchange the identical confirmation link a second time.",
        "expected": "1. Second exchange fails.\n2. User redirected to /login with notification that link was already used.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-20",
        "title": "Verification - Open-Redirect Protection on Callback",
        "steps": "1. User clicks link containing open-redirect injection: ?next=https://attacker-cebu.com.",
        "expected": "1. isSafeRedirectUrl() flags external protocol.\n2. Destination sanitized and forced to default internal path (/search or /dashboard).",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-21",
        "title": "Verification - Post-Activation Role-Based Routing",
        "steps": "1. Complete PKCE confirmation for Renter user.\n2. Complete PKCE confirmation for Landlord user.",
        "expected": "1. Renter is routed to catalog search /search.\n2. Landlord is routed to landlord KYC onboarding /landlord/onboarding.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-REG-22",
        "title": "Registration - Unactivated Account Login Prevention",
        "steps": "1. Attempt to log in with newly registered credentials before clicking email activation link.",
        "expected": "1. Login rejected with HTTP 403 AUTH_EMAIL_NOT_CONFIRMED.\n2. UI displays banner prompt to check inbox for verification email.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-23",
        "title": "Registration - Terms of Service Acceptance Requirement",
        "steps": "1. Submit registration payload with terms_accepted=false or omitted.",
        "expected": "1. Form submission rejected with HTTP 422 TERMS_REQUIRED.\n2. Inline error indicates Terms and Privacy Policy must be accepted.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-REG-24",
        "title": "Registration - Unicode & Philippine Character Support",
        "steps": "1. Register with full_name containing Philippine characters: 'José Peña Niño'.",
        "expected": "1. Registration succeeds.\n2. Name stored and retrieved with exact UTF-8 character encoding.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-REG-25",
        "title": "Registration - Concurrent Registration Race Handling",
        "steps": "1. Submit two simultaneous registration requests with identical email from two different network threads.",
        "expected": "1. Exactly one registration succeeds.\n2. Second request handled gracefully via unique constraint without crashing.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

# 2. 25 Login Test Cases (SCRUM-111)
log_cases = [
    {
        "id": "TC-LOG-01",
        "title": "Login - Valid Renter Authentication",
        "steps": "1. Submit valid Renter credentials ('mika.cebu@example.com').\n2. Call login Route Handler.",
        "expected": "1. HTTP 200 response.\n2. HTTP-only session cookies set (sb-*-auth-token, HttpOnly=true, Secure=true, SameSite=lax).\n3. Redirect target resolved to /search.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-02",
        "title": "Login - Valid Landlord Authentication",
        "steps": "1. Submit valid Landlord credentials ('carlos.landlord@example.com').",
        "expected": "1. HTTP 200 response.\n2. Session cookies set.\n3. Redirect target resolved to /landlord/dashboard.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-03",
        "title": "Login - Valid Admin Authentication",
        "steps": "1. Submit internal administrator credentials ('admin@abangcebu.internal').",
        "expected": "1. HTTP 200 response.\n2. Admin session cookies established.\n3. Redirect target resolved to /admin.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-04",
        "title": "Login - Invalid Password Rejection",
        "steps": "1. Submit valid email with incorrect password.\n2. Repeat 3 times.",
        "expected": "1. Request rejected with HTTP 401 AUTH_INVALID_CREDENTIALS.\n2. Generic message: 'Invalid email or password.'\n3. Zero session cookies set.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-05",
        "title": "Login - Non-Existent User Rejection",
        "steps": "1. Submit non-existent email address ('ghost@example.com') with arbitrary password.",
        "expected": "1. Request rejected with HTTP 401 AUTH_INVALID_CREDENTIALS.\n2. Identical error copy to TC-LOG-04 preventing user enumeration.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-06",
        "title": "Login - Unconfirmed Email Login Block",
        "steps": "1. Attempt login using 'unconfirmed.renter@example.com'.",
        "expected": "1. Request rejected with HTTP 403 AUTH_EMAIL_NOT_CONFIRMED.\n2. UI instructs user to verify email before accessing account with resend link.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-07",
        "title": "Login - Suspended User Login Interception",
        "steps": "1. Attempt login with suspended user credentials ('scammer.landlord@example.com').",
        "expected": "1. Authentication fails with HTTP 403 AUTH_ACCOUNT_SUSPENDED.\n2. Edge Middleware immediately blocks session creation.\n3. UI directs user to /suspended with support appeal instructions.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-08",
        "title": "Login - Remember Me Active (30-Day Cookie)",
        "steps": "1. Submit valid login with remember_me=true.",
        "expected": "1. Session cookies emitted with Max-Age=2592000 (30 days).\n2. Expires attribute explicitly set to 30 days in future.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-09",
        "title": "Login - Remember Me Inactive (Session Cookie)",
        "steps": "1. Submit valid login with remember_me=false.",
        "expected": "1. Session cookies emitted without explicit Max-Age.\n2. Cookies expire upon browser process termination.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-10",
        "title": "Login - Cookie Chunking (>4096 Bytes)",
        "steps": "1. Authenticate user with extensive profile metadata triggering serialized cookie payload > 4KB.",
        "expected": "1. @supabase/ssr chunks session across sb-*-auth-token.0, sb-*-auth-token.1.\n2. Middleware successfully reconstructs session without hydration drop.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-11",
        "title": "Login - Refresh Token Rotation (RTR)",
        "steps": "1. Present active refresh token to obtain fresh 1-hour JWT access token.",
        "expected": "1. New JWT issued.\n2. Old refresh token invalidated; new single-use refresh token returned.\n3. 30-second concurrency grace period active.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-12",
        "title": "Login - RTR Concurrency Grace Window",
        "steps": "1. Fire two concurrent token refresh requests within 3 seconds using the same refresh token.",
        "expected": "1. Both requests succeed within the 30-second grace window.\n2. Client receives valid rotated session without abrupt logout.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-13",
        "title": "Login - Token Family Reuse Detection & Lockout",
        "steps": "1. Malicious actor re-submits a previously rotated (revoked) refresh token after 30s grace window.",
        "expected": "1. Supabase GoTrue flags token family reuse.\n2. Entire token family instantly revoked.\n3. All active sessions for user terminated with HTTP 401 AUTH_REFRESH_TOKEN_REUSED.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-14",
        "title": "Login - Philippine Phone Number Login",
        "steps": "1. Submit phone number '09171234567' and valid password.",
        "expected": "1. Phone normalized to +639171234567.\n2. User successfully authenticated and session issued.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-15",
        "title": "Login - Password Masking & Input Type Verification",
        "steps": "1. Inspect password input element on login form.",
        "expected": "1. Field has type='password'.\n2. Toggle button reveals password in plaintext with accessible aria-label.\n3. Form auto-submits over TLS (HTTPS).",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-LOG-16",
        "title": "Login - Brute Force Rate Limiting (5 Attempts)",
        "steps": "1. Submit 5 invalid passwords consecutively from same IP within 15 minutes.",
        "expected": "1. 6th attempt rejected with HTTP 429 RATE_LIMIT_EXCEEDED.\n2. Ephemeral toast displays retry-after cooldown countdown.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-17",
        "title": "Login - Return URL Redirection (?next=)",
        "steps": "1. Access protected route /saved while unauthenticated.\n2. Redirected to /login?next=/saved.\n3. Submit valid credentials.",
        "expected": "1. After successful authentication, user is redirected to /saved instead of default role dashboard.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-18",
        "title": "Login - Open-Redirect Defense on Return URL",
        "steps": "1. Access /login?next=https://malicious-site.com.\n2. Authenticate successfully.",
        "expected": "1. isSafeRedirectUrl() flags external protocol.\n2. User redirected safely to default role landing page (/search or /dashboard).",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-19",
        "title": "Login - Session Re-Authentication Prompt",
        "steps": "1. Attempt sensitive action (e.g. payout bank details edit) after 30 minutes of idle session.",
        "expected": "1. Re-auth modal displayed.\n2. Action proceeds only upon password re-verification.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-20",
        "title": "Login - Trailing Whitespace in Login Identifier",
        "steps": "1. Enter email '  mika.cebu@example.com  ' with valid password.",
        "expected": "1. Identifier trimmed client-side and server-side.\n2. Login succeeds seamlessly.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-LOG-21",
        "title": "Login - Client State Synchronization Upon Login",
        "steps": "1. Log in on Tab A while Tab B is open on public landing page.",
        "expected": "1. Tab A logs in.\n2. Tab B detects session establishment via BroadcastChannel and updates header state to show user avatar.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-22",
        "title": "Login - Expired JWT Access Token Renewal",
        "steps": "1. Simulate expired 1-hour JWT.\n2. Next.js Server Component executes fetch request.",
        "expected": "1. Middleware transparently rotates token.\n2. Fresh JWT returned without disrupting user view.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-LOG-23",
        "title": "Login - Empty Credentials Validation",
        "steps": "1. Submit login form with empty email and password fields.",
        "expected": "1. Form displays inline validation errors for both fields.\n2. Zero network requests dispatched.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-LOG-24",
        "title": "Login - Network Disconnection During Login Request",
        "steps": "1. Simulate offline state (navigator.onLine=false) right upon submitting login credentials.",
        "expected": "1. Client displays toast: 'Network connection lost. Please check your connection and retry.'\n2. User input preserved in form fields.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-LOG-25",
        "title": "Login - Audit Logging of Successful & Failed Logins",
        "steps": "1. Execute 1 successful and 1 failed login attempt.\n2. Inspect auth audit log records.",
        "expected": "1. Both events recorded with timestamp, masked identifier, IP hash, and status.\n2. Complies with RA 10173 data privacy rules.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

# 3. 15 Password Reset Test Cases (SCRUM-112 Part 1)
rst_cases = [
    {
        "id": "TC-RST-01",
        "title": "Password Reset - Valid Recovery Request",
        "steps": "1. Submit valid email to POST /api/auth/reset-password/request.",
        "expected": "1. HTTP 200 response.\n2. Generic message: 'If an account is associated with this email, a recovery link has been dispatched.'\n3. Single-use PKCE link sent.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-02",
        "title": "Password Reset - Anti-Enumeration for Non-Existent Email",
        "steps": "1. Submit unregistered email ('nonexistent.user@example.com').",
        "expected": "1. HTTP 200 response.\n2. Message and timing match TC-RST-01 identically (artificial 200-400ms jitter applied).\n3. Zero account existence disclosure.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-03",
        "title": "Password Reset - Rate Limiting (3 in 15 mins)",
        "steps": "1. Submit 4 password reset requests within a 15-minute rolling window for the same email.",
        "expected": "1. Requests 1–3 succeed.\n2. Request 4 rejected with HTTP 429 AUTH_PASSWORD_RESET_RATE_LIMIT.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-04",
        "title": "Password Reset - Valid Recovery Link Exchange",
        "steps": "1. User clicks email link GET /auth/callback?code=RECOVERY_CODE&type=recovery.\n2. Route Handler exchanges code.",
        "expected": "1. Temporary recovery session established.\n2. Redirects to /reset-password.\n3. Session limited strictly to password update actions.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-05",
        "title": "Password Reset - Expired Recovery Token Link (>60 mins)",
        "steps": "1. User clicks recovery link after 60 minutes.",
        "expected": "1. Token exchange rejected.\n2. Redirects to /login?error=recovery_link_expired.\n3. UI displays banner prompt to request fresh reset link.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-06",
        "title": "Password Reset - Single-Use Token Replay Invalidation",
        "steps": "1. Successfully use recovery link to reach /reset-password.\n2. Attempt to open identical link again.",
        "expected": "1. Second exchange rejected.\n2. Token marked consumed.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-07",
        "title": "Password Reset - NIST Password Complexity Validation",
        "steps": "1. In /reset-password, enter weak new password ('short').\n2. Submit update.",
        "expected": "1. Rejected with HTTP 422 WEAK_PASSWORD.\n2. Inline checklist highlights missing 8-character and character-class criteria.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-08",
        "title": "Password Reset - Password Confirmation Mismatch",
        "steps": "1. Enter valid new password and non-matching confirm password.",
        "expected": "1. Client validator blocks submission.\n2. Server returns HTTP 422 PASSWORD_CONFIRMATION_MISMATCH.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-09",
        "title": "Password Reset - Successful Update & Hash Commit",
        "steps": "1. Submit valid NIST-compliant new password.",
        "expected": "1. Password hash updated in auth.users.\n2. Success response returned.\n3. User redirected to /login?message=password_reset_success.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-10",
        "title": "Password Reset - Mandatory Global Revocation",
        "steps": "1. User updates password in Browser A.\n2. Attempt API request in Browser B with pre-existing session.",
        "expected": "1. Supabase executes signOut({ scope: 'global' }).\n2. Browser B session immediately invalidated (HTTP 401).\n3. Old refresh tokens completely dead.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-11",
        "title": "Password Reset - Cookie Purging on Reset Success",
        "steps": "1. Inspect Set-Cookie headers upon password reset completion.",
        "expected": "1. All session and recovery cookies purged with Max-Age=0.\n2. Forces clean login with new credentials.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-12",
        "title": "Password Reset - Previous Password Reuse Prevention",
        "steps": "1. Attempt to set new password identical to current existing password.",
        "expected": "1. Server rejects with HTTP 422 PASSWORD_REUSE_FORBIDDEN.\n2. Prompt suggests choosing a distinct password.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-RST-13",
        "title": "Password Reset - Suspended Account Reset Block",
        "steps": "1. Request password reset for an account currently flagged profiles.is_suspended=true.",
        "expected": "1. Generic notice displayed to requester (anti-enumeration).\n2. Recovery email suppressed or contains suspension support notification.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-RST-14",
        "title": "Password Reset - Open-Redirect Guard on Recovery Callback",
        "steps": "1. Access recovery callback with ?next=https://phishing-site.com.",
        "expected": "1. isSafeRedirectUrl() strips external domain.\n2. Safely routes to /reset-password.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-RST-15",
        "title": "Password Reset - Security Audit Log Emission",
        "steps": "1. Complete successful password reset.\n2. Query auth security audit logs.",
        "expected": "1. Event PASSWORD_RESET_COMPLETED recorded with timestamp and masked identifier.\n2. Notification email sent to user confirming password change.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

# 4. 15 Sign-Out Test Cases (SCRUM-112 Part 2)
out_cases = [
    {
        "id": "TC-OUT-01",
        "title": "Sign-Out - Local Device Invalidation (scope: 'local')",
        "steps": "1. Authenticated user triggers Sign Out.\n2. Execute signOutAction({ scope: 'local' }).",
        "expected": "1. Current session deleted in auth.sessions.\n2. Set-Cookie emits Max-Age=0 and past epoch for all cookie chunks.\n3. Remote sessions on other devices remain active.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-02",
        "title": "Sign-Out - Global Multi-Device Invalidation (scope: 'global')",
        "steps": "1. User selects 'Log out of all devices'.\n2. Execute signOutAction({ scope: 'global' }).",
        "expected": "1. All rows for user_id purged from auth.sessions.\n2. All refresh token families revoked.\n3. All sessions across all devices terminated.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-03",
        "title": "Sign-Out - Multi-Chunk Cookie Zeroing (sb-*-auth-token.0..N)",
        "steps": "1. User with chunked cookies (.0, .1) triggers logout.",
        "expected": "1. Response emits Set-Cookie for base name and all chunk indices with Max-Age=0 and Expires=1970.\n2. Client cookie jar completely wiped.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-04",
        "title": "Sign-Out - BroadcastChannel Multi-Tab Synchronization",
        "steps": "1. User opens Tab A and Tab B while logged in.\n2. User clicks Sign Out in Tab A.",
        "expected": "1. Tab A executes logout and posts SIGNED_OUT to BroadcastChannel('supabase.auth.token').\n2. Tab B receives message, purges local cache, and redirects to /login?reason=signed_out.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-05",
        "title": "Sign-Out - In-Memory State & React Query Cache Flush",
        "steps": "1. Inspect TanStack Query / SWR cache and user profile store immediately following logout.",
        "expected": "1. All user query keys evicted and reset to empty state.\n2. Zero residual cached user data in browser memory.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-06",
        "title": "Sign-Out - Router Cache Purge & Back Button Guard",
        "steps": "1. User logs out from /dashboard.\n2. User clicks browser 'Back' button.",
        "expected": "1. Next.js router.refresh() purges client RSC payload cache.\n2. Protected route headers (Cache-Control: no-store) force server re-evaluation.\n3. Middleware blocks access with HTTP 307 to /login.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-07",
        "title": "Sign-Out - Offline Logout Fallback Resilience",
        "steps": "1. Disconnect network (navigator.onLine = false).\n2. User triggers Sign Out.",
        "expected": "1. Optimistic local cookie destruction executes.\n2. Memory credentials flushed.\n3. UI navigates safely to unauthenticated landing page.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-OUT-08",
        "title": "Sign-Out - Security Response Header Emission",
        "steps": "1. Inspect HTTP response headers of POST /api/auth/logout.",
        "expected": "1. Headers contain Clear-Site-Data: 'cache', 'cookies', 'storage'.\n2. Cache-Control: no-store, no-cache, must-revalidate enforced.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-OUT-09",
        "title": "Sign-Out - Modal Confirmation Flow",
        "steps": "1. User clicks 'Sign Out' in avatar menu.\n2. In confirmation modal, click 'Cancel'.",
        "expected": "1. Modal closes.\n2. User session remains fully active with zero cookie mutations.",
        "status": "Pending Dev",
        "priority": "P2 (Medium)"
    },
    {
        "id": "TC-OUT-10",
        "title": "Sign-Out - Post-Logout Redirect Destination",
        "steps": "1. Complete sign-out from /landlord/dashboard.",
        "expected": "1. User redirected to /login?reason=signed_out.\n2. Toast/Banner confirms successful logout.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-OUT-11",
        "title": "Sign-Out - Stale Token Rejection After Logout",
        "steps": "1. Capture JWT before logout.\n2. Perform logout.\n3. Attempt to use captured JWT in Authorization header for protected API.",
        "expected": "1. API rejects request with HTTP 401 Unauthorized.\n2. GoTrue session invalidation confirmed.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-12",
        "title": "Sign-Out - Concurrent Multi-Tab Sign-Out Storm",
        "steps": "1. Trigger sign-out almost simultaneously in Tab A and Tab B.",
        "expected": "1. Both tabs navigate cleanly to unauthenticated state.\n2. No uncaught race condition errors in browser console.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-OUT-13",
        "title": "Sign-Out - Session Cleanup on Suspended User Force Logout",
        "steps": "1. Admin triggers global sign-out for suspended user.",
        "expected": "1. Next request from any client of user receives session revocation.\n2. Cookies wiped and user expelled to /suspended.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-14",
        "title": "Sign-Out - Open-Redirect Defense on Post-Logout Redirect",
        "steps": "1. Call POST /api/auth/logout with redirectUrl=https://external-phish.com.",
        "expected": "1. isSafeRedirectUrl() sanitizes redirect target.\n2. Falls back to default internal /login.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-OUT-15",
        "title": "Sign-Out - Audit Log Emission for Sign-Out Event",
        "steps": "1. Execute sign-out.\n2. Check audit log entries.",
        "expected": "1. Event USER_LOGGED_OUT logged with timestamp, user_id, scope, and IP hash.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

# 5. 10 Security & Rate Limiting Test Cases
sec_cases = [
    {
        "id": "TC-SEC-01",
        "title": "Security - Login Brute Force Lockout (5 Attempts)",
        "steps": "1. Submit 5 consecutive invalid login attempts from same IP within 15 minutes.",
        "expected": "1. Attempt 6 rejected with HTTP 429 RATE_LIMIT_EXCEEDED.\n2. Temporary IP cooldown enforced.\n3. Rate limit warning displayed in UI toast.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-02",
        "title": "Security - Registration IP Rate Limiting (3 in 1hr)",
        "steps": "1. Submit 4 registrations from identical IP address within 1 hour.",
        "expected": "1. 4th registration rejected with HTTP 429 RATE_LIMIT_EXCEEDED.\n2. Bot prevention threshold enforced.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-03",
        "title": "Security - Cookie Tampering & Signature Verification",
        "steps": "1. Modify signature byte of sb-*-auth-token cookie in browser DevTools.\n2. Make request to /dashboard.",
        "expected": "1. Middleware flags corrupted signature.\n2. Cookie discarded; user redirected to /login with session invalidation.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-04",
        "title": "Security - Cross-Site Request Forgery (CSRF) Protection",
        "steps": "1. Submit cross-origin POST request to /api/auth/logout without proper Origin / Referer header.",
        "expected": "1. Request rejected with HTTP 403 CSRF_DETECTED.\n2. SameSite=Lax cookie attribute blocks cross-origin dispatch.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-05",
        "title": "Security - HTTP-Only & Secure Cookie Flags",
        "steps": "1. Inspect Set-Cookie header directives on all auth responses.",
        "expected": "1. HttpOnly=true, Secure=true, SameSite=Lax, and Path=/ present on every session token cookie.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-06",
        "title": "Security - Session Fixation Prevention",
        "steps": "1. Verify session identifier before and after successful login authentication.",
        "expected": "1. Pre-login anonymous identifier destroyed.\n2. Brand new authenticated session ID generated upon login.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-SEC-07",
        "title": "Security - Content Security Policy (CSP) Headers",
        "steps": "1. Check response headers on auth pages (/login, /register, /forgot-password).",
        "expected": "1. Strict CSP headers present, forbidding unsafe-inline scripts without nonce.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-SEC-08",
        "title": "Security - Timing Attack Resistance on Passwords",
        "steps": "1. Measure response time variance between invalid user vs invalid password on valid user.",
        "expected": "1. Response timing differences are below statistical significance (< 30ms delta).\n2. Constant-time comparison verified.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-SEC-09",
        "title": "Security - Clickjacking Frame Protection",
        "steps": "1. Attempt to embed /login inside iframe on external domain.",
        "expected": "1. X-Frame-Options: DENY and CSP frame-ancestors 'none' prevent embedding.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-SEC-10",
        "title": "Security - PII Masking in Audit Logs",
        "steps": "1. Trigger auth errors and check server log outputs.",
        "expected": "1. Passwords and tokens never logged.\n2. Emails masked (e.g. m***@example.com) for RA 10173 compliance.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

# 6. 10 Error Handling & Suspension Test Cases
err_cases = [
    {
        "id": "TC-ERR-01",
        "title": "Error Handling - Form Field Validation Errors (Inline)",
        "steps": "1. Submit invalid input fields (e.g. malformed email).\n2. Check UI surface.",
        "expected": "1. Field-specific error rendered beneath input with aria-invalid='true'.\n2. Focus moves to first invalid field.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-ERR-02",
        "title": "Error Handling - Ephemeral Rate Limit Notification (Toast)",
        "steps": "1. Trigger rate limit 429 response.",
        "expected": "1. Ephemeral floating toast appears in thumb zone.\n2. Displays cooldown timer and auto-dismisses after cooldown.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-ERR-03",
        "title": "Error Handling - Persistent Unconfirmed Email Alert (Banner)",
        "steps": "1. Attempt login with unconfirmed email.",
        "expected": "1. Docked persistent banner displays at top of login form.\n2. Contains one-click 'Resend Confirmation Email' button.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-ERR-04",
        "title": "Error Handling - Critical 500 React Error Boundary",
        "steps": "1. Trigger unhandled exception inside auth component lifecycle.",
        "expected": "1. Error boundary catches crash gracefully.\n2. Displays friendly 'Something went wrong' card with 'Reload' button.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-ERR-05",
        "title": "Error Handling - Account Suspension Enforcement",
        "steps": "1. Set profiles.is_suspended=true for active user.\n2. User clicks any navigation link.",
        "expected": "1. Middleware intercepts request with HTTP 403.\n2. Cookies zeroed immediately.\n3. User redirected to /suspended.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-ERR-06",
        "title": "Error Handling - Suspension Appeals Form",
        "steps": "1. On /suspended page, view appeal options and submit inquiry.",
        "expected": "1. Suspension reason code displayed clearly.\n2. Support email and appeal inquiry form accessible.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-ERR-07",
        "title": "Error Handling - Network Disconnection Toast Alert",
        "steps": "1. Disconnect network while interacting with auth views.",
        "expected": "1. Offline toast displayed: 'Connection lost. Offline mode active.'\n2. Reconnection automatically removes toast.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    },
    {
        "id": "TC-ERR-08",
        "title": "Error Handling - GoTrue Upstream 502/503 Outage Handling",
        "steps": "1. Simulate GoTrue gateway timeout or service unavailable.",
        "expected": "1. Route Handler catches outage.\n2. Returns HTTP 503 AUTH_SERVICE_UNAVAILABLE with friendly retry prompt.",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-ERR-09",
        "title": "Error Handling - Standardized JSON Error Contract Validation",
        "steps": "1. Call auth API routes with invalid payloads.",
        "expected": "1. All responses strictly adhere to StandardAuthError schema (code, message, status, feedbackMode, traceId).",
        "status": "Pending Dev",
        "priority": "P0 (Critical)"
    },
    {
        "id": "TC-ERR-10",
        "title": "Error Handling - Session Expiry Re-Auth Modal",
        "steps": "1. Session expires while user is typing a listing inquiry.",
        "expected": "1. Re-auth modal prompts for password without discarding draft inquiry text.",
        "status": "Pending Dev",
        "priority": "P1 (High)"
    }
]

all_100_cases = reg_cases + log_cases + rst_cases + out_cases + sec_cases + err_cases
print(f"Total Master Scenarios Compiled: {len(all_100_cases)}")
assert len(all_100_cases) == 100, f"Expected 100 scenarios, got {len(all_100_cases)}"

# 7. Update docs/testing/auth-test-plan.md
print("Updating docs/testing/auth-test-plan.md with 100-scenario matrix...")
with open('docs/testing/auth-test-plan.md', 'r') as f:
    plan_content = f.read()

# Update header ticket references to include Sprint 2 tickets
plan_content = re.sub(
    r'\*\*Jira Ticket Reference:\*\*.*?\n',
    '**Jira Ticket References:** [SCRUM-69](https://abangcebuai.atlassian.net/browse/SCRUM-69), [SCRUM-109](https://abangcebuai.atlassian.net/browse/SCRUM-109) (Master Plan), [SCRUM-110](https://abangcebuai.atlassian.net/browse/SCRUM-110) (Registration), [SCRUM-111](https://abangcebuai.atlassian.net/browse/SCRUM-111) (Login), [SCRUM-112](https://abangcebuai.atlassian.net/browse/SCRUM-112) (Reset & Sign-Out)\n',
    plan_content
)

def format_md_table(suite_title, ticket_ref, cases):
    table = f"\n### {suite_title} ({ticket_ref})\n\n"
    table += "| Test ID | Test Classification | Test Steps | Expected Confirmation / Result | Execution Status | Priority |\n"
    table += "|---|---|---|---|---|---|\n"
    for tc in cases:
        steps_clean = tc['steps'].replace('\n', '<br>')
        exp_clean = tc['expected'].replace('\n', '<br>')
        table += f"| **{tc['id']}** | {tc['title']} | {steps_clean} | {exp_clean} | {tc['status']} | {tc['priority']} |\n"
    return table

matrix_md = "## 4. Master QA Test Case Execution Matrix (100 Scenarios)\n"
matrix_md += format_md_table("Suite 1: User Registration & PKCE Verification (25 Scenarios)", "[SCRUM-110](https://abangcebuai.atlassian.net/browse/SCRUM-110)", reg_cases)
matrix_md += format_md_table("Suite 2: User Login & Session Token Lifecycle (25 Scenarios)", "[SCRUM-111](https://abangcebuai.atlassian.net/browse/SCRUM-111)", log_cases)
matrix_md += format_md_table("Suite 3: Self-Service Password Recovery (15 Scenarios)", "[SCRUM-112](https://abangcebuai.atlassian.net/browse/SCRUM-112)", rst_cases)
matrix_md += format_md_table("Suite 4: User Sign-Out & Multi-Tab Synchronization (15 Scenarios)", "[SCRUM-112](https://abangcebuai.atlassian.net/browse/SCRUM-112)", out_cases)
matrix_md += format_md_table("Suite 5: Security, Rate Limiting & Edge Defense (10 Scenarios)", "[SCRUM-109](https://abangcebuai.atlassian.net/browse/SCRUM-109)", sec_cases)
matrix_md += format_md_table("Suite 6: Error Handling & Account Suspension UX (10 Scenarios)", "[SCRUM-109](https://abangcebuai.atlassian.net/browse/SCRUM-109)", err_cases)

# Replace section 4 in auth-test-plan.md
plan_content = re.sub(
    r'## 4\. Master QA Test Case Execution Matrix.*?(## 5\. Requirements Traceability Matrix)',
    matrix_md + '\n---\n\n\\1',
    plan_content,
    flags=re.DOTALL
)

with open('docs/testing/auth-test-plan.md', 'w') as f:
    f.write(plan_content)
print("Updated docs/testing/auth-test-plan.md successfully!")

# 8. Update docs/testing/AbangCebu_Auth_Test_Specification.xlsx
print("Updating Excel workbook docs/testing/AbangCebu_Auth_Test_Specification.xlsx...")
wb = openpyxl.load_workbook('docs/testing/AbangCebu_Auth_Test_Specification.xlsx')

thin_border = Border(
    left=Side(style='thin', color='D4D4D8'),
    right=Side(style='thin', color='D4D4D8'),
    top=Side(style='thin', color='D4D4D8'),
    bottom=Side(style='thin', color='D4D4D8')
)
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=10)
align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)

def populate_sheet(sheet_name, cases, start_row=19):
    sheet = wb[sheet_name]
    # Remove merged ranges from start_row onwards
    for r in list(sheet.merged_cells.ranges):
        if r.min_row >= start_row - 1:
            sheet.unmerge_cells(str(r))
    
    # Clear cells
    for r in range(start_row, start_row + len(cases) + 30):
        for c in range(1, 13):
            cell = sheet.cell(r, c)
            if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                cell.value = None
    
    curr = start_row
    for i, tc in enumerate(cases, start=1):
        sheet.cell(curr, 2, value=i).alignment = align_center
        sheet.cell(curr, 2).font = font_bold
        sheet.cell(curr, 3, value=f"{tc['id']}: {tc['title']}").alignment = align_left
        sheet.cell(curr, 3).font = font_bold
        sheet.cell(curr, 4, value=tc['steps']).alignment = align_left
        sheet.cell(curr, 4).font = font_regular
        sheet.cell(curr, 7, value=tc['expected']).alignment = align_left
        sheet.cell(curr, 7).font = font_regular
        sheet.cell(curr, 10, value=tc['status']).alignment = align_center
        sheet.cell(curr, 10).font = font_bold
        sheet.cell(curr, 11, value=tc['priority']).alignment = align_center
        sheet.cell(curr, 11).font = font_regular
        
        for col in range(2, 12):
            sheet.cell(curr, col).border = thin_border
        curr += 1
    print(f"Populated {sheet_name} with {len(cases)} cases.")

populate_sheet('1_REGISTRATION', reg_cases, start_row=19)
populate_sheet('2_LOGIN', log_cases, start_row=19)
populate_sheet('3_LOGOUT', out_cases, start_row=19)
populate_sheet('4_PASSWORD_RESET', rst_cases, start_row=19)
populate_sheet('5_SECURITY_EDG', sec_cases + err_cases, start_row=19)

# Populate ALL_TEST_CASES
all_sheet = wb['ALL_TEST_CASES']
for r in list(all_sheet.merged_cells.ranges):
    if r.min_row >= 3:
        all_sheet.unmerge_cells(str(r))

for r in range(4, all_sheet.max_row + 150):
    for c in range(1, 13):
        cell = all_sheet.cell(r, c)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None

curr_all = 4
for i, tc in enumerate(all_100_cases, start=1):
    all_sheet.cell(curr_all, 2, value=i).alignment = align_center
    all_sheet.cell(curr_all, 2).font = font_bold
    all_sheet.cell(curr_all, 3, value=f"{tc['id']}: {tc['title']}").alignment = align_left
    all_sheet.cell(curr_all, 3).font = font_bold
    all_sheet.cell(curr_all, 4, value=tc['steps']).alignment = align_left
    all_sheet.cell(curr_all, 4).font = font_regular
    all_sheet.cell(curr_all, 7, value=tc['expected']).alignment = align_left
    all_sheet.cell(curr_all, 7).font = font_regular
    all_sheet.cell(curr_all, 10, value=tc['status']).alignment = align_center
    all_sheet.cell(curr_all, 10).font = font_bold
    all_sheet.cell(curr_all, 11, value=tc['priority']).alignment = align_center
    all_sheet.cell(curr_all, 11).font = font_regular
    
    for col in range(2, 12):
        all_sheet.cell(curr_all, col).border = thin_border
    curr_all += 1

print(f"Populated ALL_TEST_CASES with {len(all_100_cases)} cases.")
wb.save('docs/testing/AbangCebu_Auth_Test_Specification.xlsx')
print("Saved Excel workbook successfully!")

# 9. Compile Vector PDF: docs/pdf/AbangCebu_Auth_Test_Plan_Specification.pdf
print("Compiling Vector PDF for Master Test Plan...")
html_body = markdown.markdown(plan_content, extensions=['tables', 'fenced_code'])
full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @page {{
    size: A4 portrait;
    margin: 18mm 15mm 20mm 15mm;
    @bottom-right {{
      content: "Page " counter(page) " of " counter(pages);
      font-size: 8pt;
      font-family: Helvetica, sans-serif;
      color: #71717A;
    }}
    @bottom-left {{
      content: "AbangCebu AI — Master Authentication Test Plan Specification";
      font-size: 8pt;
      font-family: Helvetica, sans-serif;
      color: #71717A;
    }}
  }}
  body {{
    font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
    font-size: 9pt;
    line-height: 1.45;
    color: #18181B;
  }}
  h1 {{ font-size: 17pt; font-weight: bold; border-bottom: 2px solid #000; padding-bottom: 6px; margin-top: 0; }}
  h2 {{ font-size: 13pt; font-weight: bold; border-bottom: 1px solid #E4E4E7; padding-bottom: 4px; margin-top: 18px; }}
  h3 {{ font-size: 10.5pt; font-weight: bold; margin-top: 12px; margin-bottom: 4px; }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0;
    font-size: 8pt;
    page-break-inside: auto;
  }}
  tr {{ page-break-inside: avoid; page-break-after: auto; }}
  th, td {{
    border: 1px solid #D4D4D8;
    padding: 5px 7px;
    vertical-align: top;
  }}
  th {{
    background-color: #F4F4F5;
    font-weight: bold;
    text-align: left;
  }}
  code {{
    font-family: Monaco, Consolas, monospace;
    font-size: 7.5pt;
    background-color: #F4F4F5;
    padding: 1px 3px;
    border-radius: 2px;
  }}
  pre {{
    background-color: #F4F4F5;
    padding: 8px;
    border-radius: 3px;
    font-size: 7.5pt;
    overflow-x: auto;
  }}
  hr {{ border: none; border-top: 1px solid #E4E4E7; margin: 15px 0; }}
</style>
</head>
<body>
{html_body}
</body>
</html>
"""

weasyprint.HTML(string=full_html).write_pdf('docs/pdf/AbangCebu_Auth_Test_Plan_Specification.pdf')
pdf_size = os.path.getsize('docs/pdf/AbangCebu_Auth_Test_Plan_Specification.pdf')
print(f"Compiled Vector PDF: docs/pdf/AbangCebu_Auth_Test_Plan_Specification.pdf ({pdf_size} bytes)")

