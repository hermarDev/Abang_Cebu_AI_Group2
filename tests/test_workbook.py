#!/usr/bin/env python3
"""Assert the SCRUM-146 workbook matches the named local Supabase catalog."""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "docs/database/Database_Specification.xlsx"
TABLES = [
    "profiles", "kyc_verifications", "properties", "rental_units",
    "property_photos", "amenities", "property_amenities", "inquiries",
    "saved_listings", "reviews", "audit_logs",
]


def query(container: str, sql: str):
    result = subprocess.run(
        ["docker", "exec", container, "psql", "-U", "postgres", "-d", "postgres", "-X", "-qAt", "-c", sql],
        capture_output=True,
        text=True,
        check=True,
    )
    return json.loads(result.stdout.strip() or "null")


def table_array() -> str:
    # Fixed allowlist only; no user-controlled identifier interpolation.
    return "ARRAY[" + ",".join("'" + name + "'" for name in TABLES) + "]"


def section_rows(ws, title: str, width: int):
    title_row = next(row for row in range(1, ws.max_row + 1) if ws.cell(row, 1).value == title)
    result = []
    for row in range(title_row + 2, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, width + 1)]
        if not any(value is not None for value in values):
            break
        result.append(values)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--container",
        default=os.environ.get("SUPABASE_DB_CONTAINER", "supabase_db_scrum-143-search-db"),
        help="local Supabase Postgres container; default can be overridden by SUPABASE_DB_CONTAINER",
    )
    container = parser.parse_args().container
    names = table_array()

    with zipfile.ZipFile(WORKBOOK) as archive:
        assert archive.testzip() is None
    wb = load_workbook(WORKBOOK, read_only=False, data_only=True)
    assert len(wb.sheetnames) == 17
    assert set(TABLES).issubset(wb.sheetnames)

    columns = query(container, f"""
        SELECT coalesce(json_agg(row_to_json(q)), '[]'::json)
        FROM (
            SELECT table_name, column_name, ordinal_position, data_type, udt_name,
                   is_nullable, column_default
            FROM information_schema.columns
            WHERE table_schema='public' AND table_name=any({names})
            ORDER BY table_name, ordinal_position
        ) q
    """)
    assert len(columns) == 100, len(columns)
    for table in TABLES:
        ws = wb[table]
        expected = [item for item in columns if item["table_name"] == table]
        actual = []
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True):
            if not isinstance(row[0], int):
                continue
            assert all(value is not None for value in row[:8]), (table, row)
            actual.append((row[1], row[2], row[4], row[6]))
        catalog_values = [
            (
                item["column_name"],
                item["udt_name"] if item["data_type"] == "USER-DEFINED" else item["data_type"],
                "NO" if item["is_nullable"] == "NO" else "YES",
                item["column_default"] or "—",
            )
            for item in expected
        ]
        assert actual == catalog_values, (table, actual, catalog_values)

    indexes = query(container, f"""
        SELECT coalesce(json_agg(row_to_json(q)), '[]'::json)
        FROM (
            SELECT tablename AS table_name, indexname AS name, indexdef AS definition
            FROM pg_indexes WHERE schemaname='public' AND tablename=any({names})
            ORDER BY tablename, indexname
        ) q
    """)
    constraints = query(container, f"""
        SELECT coalesce(json_agg(row_to_json(q)), '[]'::json)
        FROM (
            SELECT c.relname AS table_name, con.conname AS name, con.contype::text AS kind,
                   pg_get_constraintdef(con.oid,true) AS definition
            FROM pg_constraint con JOIN pg_class c ON c.oid=con.conrelid
            JOIN pg_namespace n ON n.oid=c.relnamespace
            WHERE n.nspname='public' AND c.relname=any({names})
            ORDER BY c.relname, con.conname
        ) q
    """)
    assert len(indexes) == 39, len(indexes)
    assert len(constraints) == 67, len(constraints)

    for table in TABLES:
        ws = wb[table]
        expected_indexes = {
            (item["name"], item["definition"])
            for item in indexes if item["table_name"] == table
        }
        actual_indexes = {(row[0], row[1]) for row in section_rows(ws, "Indexes", 2)}
        assert actual_indexes == expected_indexes, (table, actual_indexes ^ expected_indexes)
        expected_constraints = {
            (item["name"], item["definition"])
            for item in constraints if item["table_name"] == table
        }
        actual_constraints = {(row[0], row[2]) for row in section_rows(ws, "Constraints", 3)}
        assert actual_constraints == expected_constraints, (table, actual_constraints ^ expected_constraints)

    enums = query(container, """
        SELECT coalesce(json_agg(row_to_json(q)), '[]'::json)
        FROM (
            SELECT t.typname AS name, array_agg(e.enumlabel ORDER BY e.enumsortorder) AS labels
            FROM pg_type t JOIN pg_enum e ON e.enumtypid=t.oid
            JOIN pg_namespace n ON n.oid=t.typnamespace
            WHERE n.nspname='public' GROUP BY t.typname ORDER BY t.typname
        ) q
    """)
    enum_sheet = wb["Enums & Types"]
    workbook_enums = {}
    for row in enum_sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        assert isinstance(row[0], str) and isinstance(row[2], str), row
        workbook_enums[row[0]] = row[2].split(", ")
    expected_enums = {item["name"]: item["labels"] for item in enums}
    assert workbook_enums == expected_enums, (workbook_enums, expected_enums)
    assert len(enums) == 4

    wb.close()
    print(json.dumps({
        "passed": True,
        "tables": len(TABLES),
        "columns": len(columns),
        "indexes": len(indexes),
        "constraints": len(constraints),
        "enum_types": len(enums),
        "column_metadata_match_live_catalog": True,
        "index_definitions_match_live_catalog": True,
        "constraint_definitions_match_live_catalog": True,
        "enum_labels_match_live_catalog": True,
        "xlsx_zip_integrity": "PASS",
    }, indent=2))


if __name__ == "__main__":
    main()
